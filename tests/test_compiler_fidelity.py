import tempfile
import unittest
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

from ppt_restore.rendering.compiler import NativePrimitives
from ppt_restore.rendering.geometry import CanvasSpec
from ppt_restore.rendering.scene_v2_compiler import lower_scene
from ppt_restore.contracts.schema_v2 import SceneNode, SceneSpecV2


class CompilerFidelityTests(unittest.TestCase):
    def _slide(self, elements):
        canvas = CanvasSpec()
        prs = Presentation()
        prs.slide_width, prs.slide_height = canvas.width_emu, canvas.height_emu
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        primitives = NativePrimitives()
        for element in elements:
            primitives._compile_element(slide, element, canvas)
        return prs

    def test_autoshapes_drop_theme_style_and_text_drops_autofit(self):
        prs = self._slide(
            [
                {
                    "element_type": "rounded_rect",
                    "layout_bbox_px": [10, 10, 200, 80],
                    "shape_style": {"fill": "#1C5B9E", "corner_radius_px": 8},
                },
                {
                    "element_type": "text",
                    "layout_bbox_px": [20, 100, 400, 40],
                    "text": "目标追溯",
                    "text_style": {
                        "font_size_pt": 28,
                        "font_family": "Arial",
                        "color": "#FFFFFF",
                        "word_wrap": False,
                    },
                },
            ]
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "slide.pptx"
            prs.save(path)
            with zipfile.ZipFile(path) as package:
                xml = package.read("ppt/slides/slide1.xml").decode("utf-8")
        self.assertNotIn("effectRef", xml)
        self.assertNotIn("spAutoFit", xml)
        self.assertNotIn("<p:style", xml)

    def test_ink_box_lifts_the_text_frame(self):
        canvas = CanvasSpec()
        plain = {
            "element_type": "text",
            "layout_bbox_px": [40, 80, 200, 36],
            "text": "季度",
            "text_style": {"font_size_pt": 24, "font_family": "Arial"},
        }
        ink = dict(plain)
        ink["text_style"] = dict(plain["text_style"], box_is_ink=True)
        top_plain = self._slide([plain]).slides[0].shapes[0].top
        top_ink = self._slide([ink]).slides[0].shapes[0].top
        self.assertLess(top_ink, top_plain)

    def test_home_plate_and_shape_adjust_survive_missing_enum(self):
        prs = self._slide(
            [
                {
                    "element_type": "home_plate",
                    "layout_bbox_px": [10, 10, 260, 48],
                    "shape_style": {"fill": "#1E5C9E", "shape_adjust": 0.16},
                },
                {
                    "element_type": "text",
                    "layout_bbox_px": [20, 80, 200, 30],
                    "text": "+6120W",
                    "text_style": {
                        "font_size_pt": 28,
                        "font_family": "Arial",
                        "char_spacing_pt": -1.5,
                    },
                },
            ]
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "slide.pptx"
            prs.save(path)
            with zipfile.ZipFile(path) as package:
                xml = package.read("ppt/slides/slide1.xml").decode("utf-8")
        self.assertIn('prst="homePlate"', xml)
        self.assertIn("spc=", xml)
        self.assertNotIn("effectRef", xml)

    def test_chevron_keeps_its_geometry_when_it_has_a_label(self):
        prs = self._slide(
            [
                {
                    "element_type": "chevron",
                    "layout_bbox_px": [0, 0, 220, 48],
                    "text": "Q1：基础搭建启航",
                    "shape_style": {
                        "fill": "#2E6BFF",
                        "font_size_pt": 14,
                        "color": "#FFFFFF",
                        "align": "center",
                        "vertical_anchor": "middle",
                    },
                }
            ]
        )
        shape = prs.slides[0].shapes[0]
        self.assertEqual(shape.auto_shape_type, MSO_AUTO_SHAPE_TYPE.CHEVRON)
        self.assertIn("Q1", shape.text_frame.text)

    def test_kpi_value_on_top_and_process_chevron(self):
        scene = SceneSpecV2(
            "2.0",
            "a" * 64,
            "b" * 64,
            {"width_px": 1440, "height_px": 810},
            (
                SceneNode(
                    "kpi",
                    "kpi_card",
                    "metric",
                    (100, 40, 160, 90),
                    payload={"label": "用户环比增长", "value": "15.8%"},
                    style={
                        "fill": "#2F6FED",
                        "layout": "value_on_top",
                        "value_font_size_pt": 22,
                        "label_font_size_pt": 9,
                        "value_color": "#FFFFFF",
                        "label_color": "#E8F0FF",
                    },
                ),
                SceneNode(
                    "flow",
                    "process",
                    "body",
                    (40, 200, 600, 36),
                    payload={"steps": ["目标分层落地", "差异化资源配给"]},
                    style={
                        "step_shape": "chevron",
                        "step_fills": ["#1D4ED8", "#60A5FA"],
                        "font_size_pt": 12,
                        "text_color": "#FFFFFF",
                    },
                ),
            ),
        )
        plan = lower_scene(scene)
        value = next(op for op in plan.operations if op.id == "kpi-value")
        label = next(op for op in plan.operations if op.id == "kpi-label")
        self.assertLess(value.bbox_px[1], label.bbox_px[1])
        self.assertEqual(value.style["font_size_pt"], 22)
        self.assertEqual(value.style["color"], "#FFFFFF")
        shapes = [op for op in plan.operations if op.op_type == "chevron"]
        self.assertEqual(len(shapes), 2)
        self.assertEqual(shapes[0].style["fill"], "#1D4ED8")


if __name__ == "__main__":
    unittest.main()
