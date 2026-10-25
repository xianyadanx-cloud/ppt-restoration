"""Host-agnostic defaults: body text wraps, missing z_index paints in document order."""

import tempfile
import unittest
from pathlib import Path

from ppt_restore.contracts.schema_v2 import SceneNode, SceneSpecV2
from ppt_restore.rendering.scene_v2_compiler import SceneSpecCompiler, lower_scene
from ppt_restore.rendering.text_fit import auto_word_wrap

BODY = "基于区域团队产品三线动态拆解-区域端按华东（42%）华南（28%）华中（18%）等分配，分别聚焦大客户集采，线上+线下组合渗透。"


def _scene(nodes):
    return SceneSpecV2(
        "2.0", "a" * 64, "b" * 64, {"width_px": 1440, "height_px": 810}, tuple(nodes)
    )


class AutoWrapTests(unittest.TestCase):
    def test_multiline_body_box_wraps(self):
        # 13pt ≈ 19.5px; a 3-line paragraph box 860px wide, 72px tall.
        self.assertTrue(
            auto_word_wrap(BODY, (100, 300, 860, 72), {"font_size_pt": 13}, role="body")
        )

    def test_single_line_label_never_wraps(self):
        # Too long for its box but only one line tall: overflow, never a line break.
        self.assertFalse(
            auto_word_wrap(BODY, (100, 300, 600, 26), {"font_size_pt": 13}, role="body")
        )

    def test_title_and_fitting_text_do_not_wrap(self):
        self.assertFalse(
            auto_word_wrap(
                "核心工作回顾/季度复盘/重点成果展示",
                (100, 50, 300, 200),
                {"font_size_pt": 30},
                role="title",
            )
        )
        # Fits on one line already: a tall box must not introduce a break.
        self.assertFalse(
            auto_word_wrap("15.8%", (0, 0, 200, 120), {"font_size_pt": 22}, role="metric")
        )

    def test_explicit_lines_need_extra_room_to_wrap(self):
        text = BODY[:30] + "\n" + BODY[30:60]
        self.assertFalse(
            auto_word_wrap(text, (0, 0, 300, 50), {"font_size_pt": 13}, role="body")
        )
        self.assertTrue(
            auto_word_wrap(text, (0, 0, 300, 140), {"font_size_pt": 13}, role="body")
        )

    def test_lowering_sets_wrap_only_when_absent(self):
        nodes = [
            SceneNode("body", "text", "body", (0, 0, 860, 72),
                      payload={"text": BODY}, style={"font_size_pt": 13}),
            SceneNode("pinned", "text", "body", (0, 100, 860, 72),
                      payload={"text": BODY}, style={"font_size_pt": 13, "word_wrap": False}),
            SceneNode("title", "text", "title", (0, 200, 300, 120),
                      payload={"text": BODY}, style={"font_size_pt": 30}),
        ]
        ops = {op.id: op for op in lower_scene(_scene(nodes)).operations}
        self.assertIs(ops["body"].style["word_wrap"], True)
        self.assertIs(ops["pinned"].style["word_wrap"], False)
        self.assertIs(ops["title"].style["word_wrap"], False)

    def test_compiled_body_text_frame_wraps(self):
        from pptx import Presentation

        node = SceneNode("body", "text", "body", (0, 0, 860, 72),
                         payload={"text": BODY}, style={"font_size_pt": 13})
        scene = SceneSpecV2("2.0", "a" * 64, "b" * 64,
                            {"width_px": 1440, "height_px": 810}, (node,), state="VALIDATED")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "wrap.pptx"
            SceneSpecCompiler().compile(scene, path)
            shape = Presentation(path).slides[0].shapes[0]
            self.assertTrue(shape.text_frame.word_wrap)


class PaintOrderTests(unittest.TestCase):
    def test_missing_z_index_uses_document_order_not_id(self):
        # "node-slide-plate" sorts after "node-body" alphabetically; the plate is
        # written first, so it must be painted first (behind).
        nodes = [
            SceneNode("node-slide-plate", "shape", "container", (0, 0, 1440, 810),
                      payload={"shape_type": "rect"}, style={"fill": "#FFFFFF"}),
            SceneNode("node-body", "text", "body", (10, 10, 200, 30),
                      payload={"text": "正文"}),
            SceneNode("node-a-badge", "shape", "decoration", (10, 50, 20, 20),
                      payload={"shape_type": "rect"}),
        ]
        ids = [op.id for op in lower_scene(_scene(nodes)).operations]
        self.assertEqual(ids, ["node-slide-plate", "node-body", "node-a-badge"])

    def test_explicit_z_index_still_wins(self):
        nodes = [
            SceneNode("top", "shape", "decoration", (0, 0, 10, 10),
                      payload={"shape_type": "rect"}, z_index=5),
            SceneNode("bottom", "shape", "container", (0, 0, 10, 10),
                      payload={"shape_type": "rect"}, z_index=1),
        ]
        ids = [op.id for op in lower_scene(_scene(nodes)).operations]
        self.assertEqual(ids, ["bottom", "top"])

    def test_compiled_shapes_follow_document_order(self):
        from pptx import Presentation

        nodes = (
            SceneNode("z-plate", "shape", "container", (0, 0, 1440, 810),
                      payload={"shape_type": "rect"}, style={"fill": "#FFFFFF"}),
            SceneNode("a-text", "text", "body", (10, 10, 200, 30), payload={"text": "正文"}),
        )
        scene = SceneSpecV2("2.0", "a" * 64, "b" * 64,
                            {"width_px": 1440, "height_px": 810}, nodes, state="VALIDATED")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "order.pptx"
            SceneSpecCompiler().compile(scene, path)
            shapes = list(Presentation(path).slides[0].shapes)
            self.assertEqual(shapes[-1].text_frame.text, "正文")


if __name__ == "__main__":
    unittest.main()


class GradientStopTests(unittest.TestCase):
    def test_stop_shapes_hosts_write(self):
        from ppt_restore.rendering.compiler import normalize_gradient_stops

        expected = [(0.0, "#7AAFF1"), (1.0, "#2F78BD")]
        self.assertEqual(normalize_gradient_stops([[0, "#7AAFF1"], [1, "#2F78BD"]]), expected)
        self.assertEqual(
            normalize_gradient_stops(
                [{"color": "#7AAFF1", "position": 0.0}, {"color": "#2F78BD", "position": 1.0}]
            ),
            expected,
        )
        self.assertEqual(normalize_gradient_stops([["#7AAFF1", 0], ["#2F78BD", 100]]), expected)
        self.assertEqual(normalize_gradient_stops(["#7AAFF1", "#2F78BD"]), expected)
        self.assertEqual(
            normalize_gradient_stops([{"color": "#7AAFF1", "offset": "0%"}, {"color": "#2F78BD", "offset": "100%"}]),
            expected,
        )

    def test_dict_stops_compile_without_warning_and_before_line(self):
        from pptx import Presentation

        node = SceneNode(
            "band", "shape", "decoration", (10, 10, 300, 60),
            style={
                "gradient_stops": [{"color": "#7AAFF1", "position": 0}, {"color": "#2F78BD", "position": 1}],
                "gradient_angle": 90,
                "line_color": "#FFFFFF",
                "line_width_pt": 1,
            },
        )
        scene = SceneSpecV2(
            "2.0", "a" * 64, "b" * 64, {"width_px": 1440, "height_px": 810}, (node,),
            state="VALIDATED",
        )
        with tempfile.TemporaryDirectory() as directory:
            out = Path(directory) / "g.pptx"
            result = SceneSpecCompiler().compile(scene, out)
            self.assertFalse([w for w in result.warnings if "gradient" in w], result.warnings)
            sppr = Presentation(str(out)).slides[0].shapes[0]._element.spPr
            names = [child.tag.rsplit("}", 1)[-1] for child in sppr]
            self.assertIn("gradFill", names)
            if "ln" in names:
                self.assertLess(names.index("gradFill"), names.index("ln"))


class OcrAgreementTests(unittest.TestCase):
    def test_tesseract_noise_is_not_a_conflict(self):
        from ppt_restore.pipeline.workflow import _ocr_agrees

        body = "A系列通过行业展会加KOL测评造势，首月销售额突破2600万，占总业绩38%"
        self.assertTrue(_ocr_agrees("A 系列通过行业展会加KOL测评造势，首月", [body]))
        self.assertTrue(_ocr_agrees("销售额突破2600万，占总业绩 38%", [body]))
        self.assertTrue(_ocr_agrees('端通过"技术套餐+ 性价', ["产品端通过“技术套餐+性价比捆绑”适配区域需求"]))
        # dropped characters on a tinted plate
        self.assertTrue(_ocr_agrees("户集，线上 + 线下组合渗透。团", ["分别聚焦大客户集采，线上+线下组合渗透。团队端"]))

    def test_wrong_numbers_and_wrong_text_still_conflict(self):
        from ppt_restore.pipeline.workflow import _ocr_agrees

        self.assertFalse(_ocr_agrees("销售额突破2600万", ["销售额突破2800万"]))
        self.assertFalse(_ocr_agrees("15.8%", ["15.3%"]))
        self.assertFalse(_ocr_agrees("核心市场提频", ["新品引爆市场"]))


class RunsTextConsistencyTests(unittest.TestCase):
    def test_runs_that_differ_from_text_need_review(self):
        from ppt_restore.pipeline.workflow import validate_scene_proposal

        class Evidence:
            ocr_lines = ()
            regions = ()
            palette = ()
            guides = ()
            visual_primitives = ()

        good = SceneNode("a", "text", "label", (0, 0, 200, 30), payload={
            "text": "01.内容价值深化", "runs": [{"text": "01."}, {"text": "内容价值深化"}]},
            evidence_refs=("r",))
        bad = SceneNode("b", "text", "label", (0, 40, 200, 30), payload={
            "text": "01.内容价值深化", "runs": [{"text": "01. 内"}, {"text": "容价值深化"}]},
            evidence_refs=("r",))
        _, report = validate_scene_proposal(_scene([good, bad]), Evidence())
        joined = " ".join(report["review_reasons"])
        self.assertIn("b: payload.runs spell", joined)
        self.assertNotIn("a: payload.runs", joined)
