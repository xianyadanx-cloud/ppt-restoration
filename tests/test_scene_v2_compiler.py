import tempfile
import unittest
import zipfile
from dataclasses import replace
from pathlib import Path

from PIL import Image

from ppt_restore.contracts.schema_v2 import SceneNode, SceneSpecV2
from ppt_restore.platform.evidence import extract_evidence
from ppt_restore.platform.provenance import sha256_file
from ppt_restore.rendering.scene_v2_compiler import SceneSpecCompiler, lower_scene


class SceneV2CompilerTests(unittest.TestCase):
    def test_native_table_progress_reflows_with_cell_dimensions(self):
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as directory:
            case = Path(directory)
            scene = self._scene(case)
            payload = {
                "cells": [["任务", "完成"], ["A", "50%"]],
                "merges": [[0, 0, 0, 1]],
                "column_widths": [1, 1],
                "row_heights": [1, 1],
                "progress_bars": [
                    {
                        "row": 1,
                        "column": 1,
                        "value": 0.5,
                        "marker": True,
                        "rounded": True,
                        "gradient_stops": [[0, "#9E2E25"], [1, "#E17A72"]],
                        "center_y_fraction": 0.8,
                    }
                ],
            }
            node = SceneNode(
                "table", "table", "body", (0, 0, 800, 400), payload=payload
            )
            scene = replace(scene, nodes=(node,), state="VALIDATED")

            def build(name, current):
                path = case / name
                SceneSpecCompiler().compile(current, path, strict_native=True)
                return Presentation(path).slides[0].shapes

            before = build("before.pptx", scene)
            self.assertTrue(before[0].has_table)
            self.assertTrue(before[0].table.cell(0, 0).is_merge_origin)
            self.assertIn("完成", before[0].table.cell(0, 0).text)
            self.assertEqual(before[0].table.cell(1, 1).text, "50%")
            marker = next(s for s in before if s.name.endswith(":marker"))
            self.assertEqual(marker.width, marker.height)
            from pptx.oxml.ns import qn

            fill = next(s for s in before if s.name.endswith(":fill"))
            self.assertIsNotNone(fill._element.spPr.find(qn("a:gradFill")))
            updated = dict(payload, column_widths=[3, 1], row_heights=[3, 1])
            after = build(
                "after.pptx", replace(scene, nodes=(replace(node, payload=updated),))
            )
            track_before = next(s for s in before if s.name.endswith(":track"))
            track_after = next(s for s in after if s.name.endswith(":track"))
            self.assertGreater(track_after.left, track_before.left)
            self.assertGreater(track_after.top, track_before.top)
            self.assertAlmostEqual(
                track_after.width / track_before.width, 0.5, places=5
            )
            self.assertEqual(after[0].table.cell(1, 1).text, "50%")

    def test_partial_build_inventory_excludes_future_blocks(self):
        from ppt_restore.cli.services import _reconcile_content_after_build
        from ppt_restore.platform.io import write_json
        from ppt_restore.quality.content_inventory import scene_to_content_inventory

        with tempfile.TemporaryDirectory() as directory:
            case = Path(directory)
            scene = self._scene(case)
            scene = replace(
                scene,
                nodes=(
                    replace(scene.nodes[1], block_id="b1"),
                    replace(scene.nodes[2], block_id="b2"),
                ),
                blocks=(
                    {"id": "b1", "bbox_norm": [0, 0, 500, 1000]},
                    {"id": "b2", "bbox_norm": [500, 0, 500, 1000]},
                ),
                state="VALIDATED",
            )
            (case / "scene.json").write_text(scene.to_json(), encoding="utf-8")
            write_json(
                case / "content_inventory.json", scene_to_content_inventory(scene)
            )
            output = case / "partial.pptx"
            result = SceneSpecCompiler().compile(scene, output, blocks=["b1"])
            partial = _reconcile_content_after_build(
                result, case, output, blocks=["b1"]
            )
            self.assertTrue(partial["valid"])
            self.assertFalse(partial["scope"]["full_slide"])
            full = _reconcile_content_after_build(result, case, output)
            self.assertFalse(full["valid"])

    def _scene(self, case):
        canonical = case / "canonical.png"
        Image.new("RGB", (1440, 810), "white").save(canonical)
        extract_evidence(canonical, ocr_mode="none", output_path=case / "evidence.json")
        nodes = (
            SceneNode(
                "background",
                "shape",
                "container",
                (0, 0, 1440, 810),
                z_index=0,
                style={"fill": "#FFFFFF"},
            ),
            SceneNode(
                "title",
                "text",
                "title",
                (40, 30, 400, 50),
                z_index=1,
                payload={"text": "Editable"},
                style={"font_size_pt": 24, "weight": "bold", "color": "#000000"},
            ),
            SceneNode(
                "kpi",
                "kpi_card",
                "metric",
                (500, 40, 180, 140),
                z_index=2,
                payload={"label": "规模", "value": "38万", "delta": "缺口16%"},
                style={"fill": "#DCEBFF"},
            ),
            SceneNode(
                "progress",
                "progress_bar",
                "metric",
                (40, 120, 300, 20),
                z_index=3,
                payload={"value": 0.8, "min": 0, "max": 1, "label": "80%"},
                style={"fill": "#2563EB"},
            ),
        )
        return SceneSpecV2(
            "2.0",
            sha256_file(canonical),
            sha256_file(case / "evidence.json"),
            {"width_px": 1440, "height_px": 810},
            nodes,
        )

    def test_lower_scene_emits_deterministic_operations(self):
        with tempfile.TemporaryDirectory() as directory:
            scene = self._scene(Path(directory))
            plan = lower_scene(scene)
            self.assertGreaterEqual(len(plan.operations), 7)
            self.assertEqual(
                [op.id for op in plan.operations],
                [op.id for op in lower_scene(scene).operations],
            )
            self.assertTrue(any(op.op_type == "text" for op in plan.operations))

    def test_compiler_writes_native_pptx_and_manifest(self):
        with tempfile.TemporaryDirectory() as directory:
            case = Path(directory)
            scene = self._scene(case)
            scene = replace(scene, state="VALIDATED")
            # ingest normally writes scene.json; this test supplies the same
            # validated bytes to exercise the compiler in isolation.
            (case / "scene.json").write_text(
                scene.to_json(indent=2) + "\n", encoding="utf-8"
            )
            output = case / "result.pptx"
            result = SceneSpecCompiler().compile(
                scene, output, canonical_path=case / "canonical.png", case_dir=case
            )
            self.assertTrue(output.is_file())
            self.assertGreater(result.element_count, 0)
            self.assertTrue((case / "render_plan.json").is_file())
            self.assertTrue(list((case / "builds").glob("*.json")))
            self.assertEqual(result.picture_count, 0)

    def test_hybrid_image_is_allowed_but_strict_native_rejects_it(self):
        with tempfile.TemporaryDirectory() as directory:
            case = Path(directory)
            scene = self._scene(case)
            image_node = SceneNode(
                "ornament",
                "image",
                "decoration",
                (900, 40, 120, 100),
                payload={"source_bbox_norm": [0, 0, 100, 100]},
                render_strategy="raster",
            )
            scene = replace(scene, nodes=scene.nodes + (image_node,), state="VALIDATED")
            (case / "scene.json").write_text(scene.to_json(), encoding="utf-8")
            result = SceneSpecCompiler().compile(
                scene,
                case / "hybrid.pptx",
                canonical_path=case / "canonical.png",
                case_dir=case,
            )
            self.assertEqual(result.picture_count, 1)
            with self.assertRaises(ValueError):
                SceneSpecCompiler().compile(
                    scene,
                    case / "strict.pptx",
                    canonical_path=case / "canonical.png",
                    case_dir=case,
                    strict_native=True,
                )

    def test_v2_block_filter_and_progress_marker(self):
        with tempfile.TemporaryDirectory() as directory:
            case = Path(directory)
            base = self._scene(case)
            blocks = (
                {"id": "b1", "bbox_norm": [0, 0, 500, 1000]},
                {"id": "b2", "bbox_norm": [500, 0, 500, 1000]},
            )
            nodes = (
                replace(base.nodes[0], block_id=None),
                replace(base.nodes[1], block_id="b1"),
                replace(base.nodes[2], block_id="b2"),
                replace(
                    base.nodes[3],
                    block_id="b1",
                    payload={"value": 0.8, "marker": True},
                    style={
                        "fill": "#C2410C",
                        "gradient_stops": [[0, "#9A3412"], [1, "#FB7185"]],
                        "marker_shadow": {
                            "color": "#000000",
                            "opacity": 0.2,
                            "blur_px": 3,
                            "distance_px": 1,
                        },
                    },
                ),
            )
            scene = replace(base, nodes=nodes, blocks=blocks, state="VALIDATED")
            plan = lower_scene(scene, blocks=["b1"])
            ids = {op.id for op in plan.operations}
            self.assertIn("title", ids)
            self.assertIn("progress-marker", ids)
            self.assertNotIn("kpi-base", ids)
            result = SceneSpecCompiler().compile(
                scene,
                case / "b1.pptx",
                canonical_path=case / "canonical.png",
                blocks=["b1"],
            )
            self.assertEqual(result.picture_count, 0)
            with zipfile.ZipFile(case / "b1.pptx") as package:
                slide_xml = package.read("ppt/slides/slide1.xml").decode("utf-8")
            self.assertIn("<a:gradFill", slide_xml)
            self.assertIn("<a:outerShdw", slide_xml)
            self.assertGreaterEqual(slide_xml.count("<p:sp>"), 5)


if __name__ == "__main__":
    unittest.main()
