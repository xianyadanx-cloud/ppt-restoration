"""CLI handoff test with a simulated renderer, NOT visual acceptance evidence."""

import contextlib
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image, ImageDraw
from pptx import Presentation

from ppt_restore.cli import main
from ppt_restore.contracts.schema_v2 import SceneNode, SceneSpecV2
from ppt_restore.pipeline.block_gate import init_block_gate
from ppt_restore.pipeline.host_tasks import next_task
from ppt_restore.pipeline.workflow import ingest_scene, prepare_case
from ppt_restore.platform.evidence import extract_evidence
from ppt_restore.platform.ocr import ImportedTextProvider, TextDetection
from ppt_restore.platform.provenance import sha256_file
from ppt_restore.quality.audit import make_content_audit
from ppt_restore.rendering.renderers import RenderResult


class SimulatedGeometryBackend:
    name = "libreoffice"
    available = True

    def render(self, pptx, output):
        # Exercise actual PPT compilation and scoring, but simulate its pixels.
        output = Path(output)
        output.mkdir(parents=True, exist_ok=True)
        deck = Presentation(pptx)
        image = Image.new("RGB", (1440, 810), "white")
        draw = ImageDraw.Draw(image)
        for shape in deck.slides[0].shapes:
            if shape.has_text_frame:
                x = round(shape.left / deck.slide_width * 1440)
                y = round(shape.top / deck.slide_height * 810)
                w = round(shape.width / deck.slide_width * 1440)
                h = round(shape.height / deck.slide_height * 810)
                draw.rectangle((x, y, x + w - 1, y + h - 1), fill="black")
        path = output / "simulated.png"
        image.save(path)
        return RenderResult(
            self.name, True, images=[str(path)], pptx_sha256=sha256_file(pptx)
        )


class OptimizerWorkflowTests(unittest.TestCase):
    def test_build_optimize_recover_audit_ingest_rebuild(self):
        with (
            tempfile.TemporaryDirectory() as directory,
            contextlib.redirect_stdout(io.StringIO()),
        ):
            root = Path(directory)
            source = root / "source.png"
            image = Image.new("RGB", (1440, 810), "white")
            ImageDraw.Draw(image).rectangle((26, 10, 125, 39), fill="black")
            image.save(source)
            case = root / "case"
            prepare_case(source, case, ocr_mode="none")
            extract_evidence(
                case / "canonical.png",
                output_path=case / "evidence.json",
                ocr_provider=ImportedTextProvider(
                    [TextDetection("Hello", (26, 10, 100, 30), 0.99, "imported")]
                ),
            )
            scene = SceneSpecV2(
                "2.0",
                sha256_file(case / "canonical.png"),
                sha256_file(case / "evidence.json"),
                {"width_px": 1440, "height_px": 810},
                (
                    SceneNode(
                        "title",
                        "text",
                        "title",
                        (10, 10, 100, 30),
                        block_id="header",
                        evidence_refs=("ocr-0001",),
                        payload={"text": "Hello"},
                    ),
                ),
                blocks=({"id": "header", "bbox_norm": [0, 0, 1000, 1000]},),
            )
            proposal = case / "scene.proposed.json"
            proposal.write_text(scene.to_json(), encoding="utf-8")
            audit = case / "content_audit.proposed.json"
            audit.write_text(
                make_content_audit(scene, confirmed=True).to_json(), encoding="utf-8"
            )
            self.assertEqual(
                ingest_scene(case, proposal, content_audit_path=audit)["status"],
                "VALIDATED",
            )
            init_block_gate(
                case, proposal
            )  # Synthetic fixture, not a real user approval.
            self.assertEqual(main(next_task(case)["command_argv"][1:]), 0)
            before = sha256_file(case / "scene.json")
            with (
                patch(
                    "ppt_restore.cli.commands.optimize._backend",
                    return_value=SimulatedGeometryBackend(),
                ),
                patch(
                    "ppt_restore.quality.render_probe.run_render_probe",
                    return_value={"automatic_optimization_allowed": True},
                ),
                patch(
                    "ppt_restore.quality.cache_context.visual_cache_context",
                    return_value={"test": "simulated"},
                ),
            ):
                self.assertEqual(
                    main(
                        [
                            "optimize",
                            str(case),
                            "--block",
                            "header",
                            "--max-rounds",
                            "1",
                            "--max-candidates",
                            "2",
                        ]
                    ),
                    0,
                )
            self.assertEqual(sha256_file(case / "scene.json"), before)
            task = next_task(case)
            self.assertEqual(task["status"], "CANDIDATE_AUDIT_REQUIRED")
            candidate = SceneSpecV2.from_json(
                Path(task["read_files"][1]).read_text(encoding="utf-8")
            )
            self.assertEqual(candidate.nodes[0].payload["text"], "Hello")
            Path(task["required_outputs"][0]).write_text(
                make_content_audit(candidate, confirmed=True).to_json(),
                encoding="utf-8",
            )
            self.assertEqual(main(task["after_audit_argv"][1:]), 0)
            task = next_task(case)
            self.assertEqual(task["status"], "BUILD_REQUIRED")
            self.assertEqual(main(task["command_argv"][1:]), 0)
            self.assertEqual(next_task(case)["status"], "REVIEW_REQUIRED")
