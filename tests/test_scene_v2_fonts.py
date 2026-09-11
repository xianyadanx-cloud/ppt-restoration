import io
import unittest
import zipfile

from pptx import Presentation
from pptx.oxml.ns import qn

from ppt_restore.rendering.compiler import set_paragraph_spacing, set_run_typeface


class SceneV2FontTests(unittest.TestCase):
    def test_spacing_units_and_invalid_values(self):
        prs = Presentation()
        paragraph = (
            prs.slides.add_slide(prs.slide_layouts[6])
            .shapes.add_textbox(0, 0, 1000000, 300000)
            .text_frame.paragraphs[0]
        )
        set_paragraph_spacing(paragraph, {"line_spacing": 1.25, "space_after_pt": 4})
        self.assertEqual(paragraph.line_spacing, 1.25)
        self.assertEqual(paragraph.space_after.pt, 4)
        set_paragraph_spacing(paragraph, {"line_spacing_pt": 16})
        self.assertEqual(paragraph.line_spacing.pt, 16)
        for value in (0, -1, float("nan"), float("inf")):
            with self.assertRaises(ValueError):
                set_paragraph_spacing(paragraph, {"line_spacing": value})

    def test_run_records_explicit_east_asian_typeface(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(0, 0, 1000000, 300000)
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = "季度工作攻坚策略"
        set_run_typeface(run, "PingFang SC")
        self.assertEqual(run._r.get_or_add_rPr().get("lang"), "zh-CN")
        self.assertIsNone(run._r.get_or_add_rPr().get(qn("a:lang")))

        payload = io.BytesIO()
        prs.save(payload)
        with zipfile.ZipFile(io.BytesIO(payload.getvalue())) as archive:
            xml = archive.read("ppt/slides/slide1.xml").decode("utf-8")

        self.assertIn('a:ea typeface="PingFang SC"', xml)
        self.assertIn('lang="zh-CN"', xml)


if __name__ == "__main__":
    unittest.main()
