"""High-level Pythonic Native PPTX Builder with 0-1000 Normalized Coordinates.

Designed specifically for AI Coding Agents to construct high-fidelity,
fully native, editable PowerPoint presentations without boilerplate code.
"""

from __future__ import annotations
import os
from typing import Any, Dict, List, Optional, Tuple, Union
from PIL import Image

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls
except ImportError:
    # Allow importing module even before pptx is installed; functions will check runtime
    Presentation = None


def color_to_hex(val: Any) -> str:
    """Convert color input to 6-character hex string (RRGGBB)."""
    if isinstance(val, RGBColor):
        return f"{val[0]:02X}{val[1]:02X}{val[2]:02X}"
    if isinstance(val, (tuple, list)) and len(val) == 3:
        return f"{int(val[0]):02X}{int(val[1]):02X}{int(val[2]):02X}"
    if isinstance(val, str):
        val = val.strip().lstrip("#")
        named_colors = {
            "white": "FFFFFF",
            "black": "000000",
            "gray": "64748B",
            "light_gray": "F1F5F9",
            "slate": "1E293B",
            "blue": "2563EB",
            "dark_blue": "0F172A",
            "indigo": "4F46E5",
            "cyan": "06B6D4",
            "emerald": "10B981",
            "amber": "F59E0B",
            "rose": "F43F5E",
            "purple": "7C3AED",
        }
        return named_colors.get(val.lower(), val).upper()
    return "334155"


class Tokens:
    """Standardized Design Tokens for typography scale, spacing, and radius."""

    # Typography Scales (in pt)
    FONT_HERO = 36.0
    FONT_TITLE_LG = 32.0
    FONT_TITLE_MD = 24.0
    FONT_KPI_VAL = 22.0
    FONT_SECTION_H2 = 16.0
    FONT_CARD_TITLE = 14.0
    FONT_BODY = 11.0
    FONT_BODY_SM = 10.0
    FONT_BADGE = 8.5
    FONT_CAPTION = 7.5

    # Normalized Spacings (0-1000 scale)
    GAP_XS = 4.0
    GAP_SM = 8.0
    GAP_MD = 14.0
    GAP_LG = 20.0
    GAP_XL = 30.0

    # Margins & Paddings
    PAD_CARD = 12.0
    PAD_BADGE = 6.0


def estimate_text_width_pt(text: Any, font_size_pt: float) -> float:
    """Estimate rendered text width in points (pt) based on typographic character weights."""
    if not text:
        return 0.0
    width = 0.0
    for ch in str(text):
        if ord(ch) > 127:
            # CJK full-width characters (Chinese, Japanese, full-width punctuation) ~ 1.0 em
            width += font_size_pt * 1.0
        elif ch in ".,:;!'iIl ":
            # Narrow ASCII characters ~ 0.28 em
            width += font_size_pt * 0.28
        elif ch in "mwMW@#%&":
            # Wide ASCII characters ~ 0.85 em
            width += font_size_pt * 0.85
        elif ch.isupper() or ch.isdigit():
            # Uppercase letters & standard digits ~ 0.62 em
            width += font_size_pt * 0.62
        else:
            # Standard lowercase ASCII ~ 0.52 em
            width += font_size_pt * 0.52
    return width


def calculate_safe_font_size(
    text: Any,
    target_width_pt: float,
    desired_font_size: float,
    min_font_size: float = 7.0,
    max_lines: int = 1,
) -> float:
    """Calculate maximum safe font size that fits within target width without accidental wrapping."""
    if max_lines > 1 or target_width_pt <= 0 or not text:
        return float(desired_font_size)

    current_size = float(desired_font_size)
    # Available usable width with 8% safety buffer for font kerning/internal padding
    safe_available = target_width_pt * 0.92

    while current_size > min_font_size:
        est_w = estimate_text_width_pt(text, current_size)
        if est_w <= safe_available:
            break
        current_size -= 0.5

    return max(min_font_size, round(current_size, 1))


def apply_gradient_fill(
    element_spPr,
    start_color: str,
    end_color: str,
    angle: float = 90.0,
    stops: Optional[List[Tuple[float, str]]] = None,
) -> None:
    """Inject native OpenXML gradient fill (<a:gradFill>) into a shape or background."""
    if parse_xml is None:
        return

    # Clear existing fills
    for child in list(element_spPr):
        tag_name = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag_name in ("solidFill", "noFill", "gradFill", "blipFill", "pattFill"):
            element_spPr.remove(child)

    if stops:
        gs_items = ""
        for pos_ratio, color_val in stops:
            pos_val = int(pos_ratio * 100000)
            hex_c = color_to_hex(color_val)
            gs_items += f'<a:gs pos="{pos_val}"><a:srgbClr val="{hex_c}"/></a:gs>'
    else:
        c1 = color_to_hex(start_color)
        c2 = color_to_hex(end_color)
        gs_items = (
            f'<a:gs pos="0"><a:srgbClr val="{c1}"/></a:gs>'
            f'<a:gs pos="100000"><a:srgbClr val="{c2}"/></a:gs>'
        )

    angle_val = int(angle * 60000)  # OpenXML angle is 1/60000 of a degree
    grad_xml = (
        f'<a:gradFill {nsdecls("a")} flip="none" rotWithShape="1">'
        f'  <a:gsLst>{gs_items}</a:gsLst>'
        f'  <a:lin ang="{angle_val}" scaled="0"/>'
        f'</a:gradFill>'
    )
    grad_elem = parse_xml(grad_xml)
    element_spPr.append(grad_elem)


def parse_color(val: Any) -> Optional[RGBColor]:
    """Parse color into pptx RGBColor. Supports hex, RGB tuple, and named colors."""
    if val is None or val == "transparent" or val == "none":
        return None
    if isinstance(val, RGBColor):
        return val
    if isinstance(val, (tuple, list)) and len(val) == 3:
        return RGBColor(int(val[0]), int(val[1]), int(val[2]))
    if isinstance(val, str):
        val = val.strip().lstrip("#")
        named_colors = {
            "white": "FFFFFF",
            "black": "000000",
            "gray": "64748B",
            "light_gray": "F1F5F9",
            "slate": "1E293B",
            "blue": "2563EB",
            "dark_blue": "0F172A",
            "indigo": "4F46E5",
            "cyan": "06B6D4",
            "emerald": "10B981",
            "amber": "F59E0B",
            "rose": "F43F5E",
            "purple": "7C3AED",
        }
        hex_str = named_colors.get(val.lower(), val)
        if len(hex_str) == 6:
            r = int(hex_str[0:2], 16)
            g = int(hex_str[2:4], 16)
            b = int(hex_str[4:6], 16)
            return RGBColor(r, g, b)
    return RGBColor(50, 50, 50)


def normalize_box(
    box: Optional[Union[List[float], Tuple[float, ...]]] = None,
    left: Optional[float] = None,
    top: Optional[float] = None,
    width: Optional[float] = None,
    height: Optional[float] = None,
) -> Tuple[float, float, float, float]:
    """Normalize box coordinates into (left, top, width, height) in 0-1000 scale.
    Clamps bounds and prevents overflow gracefully.
    """
    if box is not None and len(box) == 4:
        l, t, w, h = float(box[0]), float(box[1]), float(box[2]), float(box[3])
    else:
        l = float(left if left is not None else 0)
        t = float(top if top is not None else 0)
        w = float(width if width is not None else 100)
        h = float(height if height is not None else 100)

    # Clamping & Sanity Checks
    l = max(0.0, min(1000.0, l))
    t = max(0.0, min(1000.0, t))
    w = max(1.0, min(1000.0 - l, w))
    h = max(1.0, min(1000.0 - t, h))
    return l, t, w, h


class SlideBuilder:
    """High-level builder for constructing native editable PowerPoint slides."""

    def __init__(
        self,
        aspect_ratio: str = "16:9",
        width_inch: Optional[float] = None,
        height_inch: Optional[float] = None,
        bg_color: Optional[str] = "#FFFFFF",
    ):
        if Presentation is None:
            raise RuntimeError(
                "python-pptx is not installed. Please run: pip install python-pptx"
            )

        self.prs = Presentation()
        # Set dimensions
        if width_inch and height_inch:
            self.width_inch = width_inch
            self.height_inch = height_inch
        elif aspect_ratio == "4:3":
            self.width_inch = 10.0
            self.height_inch = 7.5
        else:  # Default 16:9 widescreen
            self.width_inch = 13.333
            self.height_inch = 7.5

        self.prs.slide_width = Inches(self.width_inch)
        self.prs.slide_height = Inches(self.height_inch)

        # Blank layout
        blank_slide_layout = self.prs.slide_layouts[6]
        self.slide = self.prs.slides.add_slide(blank_slide_layout)

        if bg_color:
            self.set_background(color=bg_color)

    @classmethod
    def from_image(
        cls,
        image_path: str,
        bg_color: Optional[str] = "#FFFFFF",
    ) -> "SlideBuilder":
        """Automatically detect aspect ratio and dimensions from an input image."""
        if not os.path.exists(image_path):
            return cls(aspect_ratio="16:9", bg_color=bg_color)

        with Image.open(image_path) as img:
            img_w, img_h = img.size

        ratio = img_w / float(img_h) if img_h > 0 else 16.0 / 9.0
        # Standardize height to 7.5 inches and scale width
        h_inch = 7.5
        w_inch = round(h_inch * ratio, 3)

        aspect_str = "16:9"
        if abs(ratio - (4.0 / 3.0)) < 0.08:
            aspect_str = "4:3"
            w_inch, h_inch = 10.0, 7.5
        elif abs(ratio - (16.0 / 9.0)) < 0.08:
            aspect_str = "16:9"
            w_inch, h_inch = 13.333, 7.5

        return cls(
            aspect_ratio=aspect_str,
            width_inch=w_inch,
            height_inch=h_inch,
            bg_color=bg_color,
        )

    def _to_emu(
        self,
        box: Optional[Union[List[float], Tuple[float, ...]]] = None,
        left: Optional[float] = None,
        top: Optional[float] = None,
        width: Optional[float] = None,
        height: Optional[float] = None,
    ) -> Tuple[Any, Any, Any, Any]:
        """Convert 0-1000 coordinates to pptx EMU inches."""
        l, t, w, h = normalize_box(box, left, top, width, height)
        return (
            Inches(l * self.width_inch / 1000.0),
            Inches(t * self.height_inch / 1000.0),
            Inches(w * self.width_inch / 1000.0),
            Inches(h * self.height_inch / 1000.0),
        )

    def set_background(
        self,
        color: Optional[str] = "#FFFFFF",
        image_path: Optional[str] = None,
        gradient_colors: Optional[List[str]] = None,
        gradient_angle: float = 90.0,
    ) -> None:
        """Set solid background color, full-bleed background image, or native gradient fill."""
        if image_path and os.path.exists(image_path):
            self.slide.shapes.add_picture(
                image_path, Inches(0), Inches(0), Inches(self.width_inch), Inches(self.height_inch)
            )
            return

        if gradient_colors and len(gradient_colors) >= 2:
            background = self.slide.background
            # Get background spPr / bgPr
            bgPr = background._element.xpath(".//p:bgPr")
            if bgPr:
                apply_gradient_fill(bgPr[0], gradient_colors[0], gradient_colors[1], angle=gradient_angle)
            return

        if color:
            bg_rgb = parse_color(color)
            if bg_rgb:
                background = self.slide.background
                fill = background.fill
                fill.solid()
                fill.fore_color.rgb = bg_rgb

    def add_header(
        self,
        title: str,
        subtitle: Optional[str] = None,
        category_tag: Optional[str] = None,
        box: Optional[List[float]] = None,
        title_size: int = 24,
        subtitle_size: int = 12,
        title_color: str = "#0F172A",
        subtitle_color: str = "#64748B",
        tag_color: str = "#2563EB",
        tag_bg: str = "#EFF6FF",
    ) -> None:
        """Add a professional consulting-style title header with optional category tag and subtitle."""
        l, t, w, h = self._to_emu(box or [40, 35, 920, 100])
        tx_box = self.slide.shapes.add_textbox(l, t, w, h)
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.space_after = Pt(4)

        if category_tag:
            run_tag = p.add_run()
            run_tag.text = f"[{category_tag.upper()}]  "
            run_tag.font.size = Pt(max(9, title_size - 14))
            run_tag.font.bold = True
            run_tag.font.color.rgb = parse_color(tag_color)

        run_title = p.add_run()
        run_title.text = title
        run_title.font.size = Pt(title_size)
        run_title.font.bold = True
        run_title.font.color.rgb = parse_color(title_color)

        if subtitle:
            p_sub = tf.add_paragraph()
            p_sub.space_before = Pt(2)
            run_sub = p_sub.add_run()
            run_sub.text = subtitle
            run_sub.font.size = Pt(subtitle_size)
            run_sub.font.color.rgb = parse_color(subtitle_color)

    def add_card(
        self,
        box: List[float],
        bg_color: Optional[str] = "#FFFFFF",
        border_color: Optional[str] = "#E2E8F0",
        border_width_pt: float = 1.0,
        radius: bool = True,
        gradient_colors: Optional[List[str]] = None,
        gradient_angle: float = 90.0,
        rotation: float = 0.0,
    ) -> Any:
        """Add a native background card container (solid or native gradient fill)."""
        l, t, w, h = self._to_emu(box)
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = self.slide.shapes.add_shape(shape_type, l, t, w, h)
        if rotation != 0.0:
            shape.rotation = rotation

        if gradient_colors and len(gradient_colors) >= 2:
            apply_gradient_fill(shape._element.spPr, gradient_colors[0], gradient_colors[1], angle=gradient_angle)
        elif bg_color and bg_color.lower() != "transparent":
            shape.fill.solid()
            shape.fill.fore_color.rgb = parse_color(bg_color)
        else:
            shape.fill.background()

        if border_color and border_color.lower() != "transparent":
            shape.line.color.rgb = parse_color(border_color)
            shape.line.width = Pt(border_width_pt)
        else:
            shape.line.fill.background()

        return shape

    def add_polygon(
        self,
        points: List[Tuple[float, float]],
        bg_color: Optional[str] = "#FFFFFF",
        border_color: Optional[str] = None,
        border_width_pt: float = 1.0,
        gradient_colors: Optional[List[str]] = None,
        gradient_angle: float = 90.0,
        rotation: float = 0.0,
    ) -> Any:
        """Add a native freeform polygon from a list of (x, y) points in 0-1000 coordinates."""
        if not points or len(points) < 3:
            raise ValueError("Polygon requires at least 3 points.")

        # Convert first point to EMU
        p0_x = Inches(points[0][0] * self.width_inch / 1000.0)
        p0_y = Inches(points[0][1] * self.height_inch / 1000.0)
        ff_builder = self.slide.shapes.build_freeform(p0_x, p0_y)

        # Convert remaining points to EMU
        emu_points = [
            (
                Inches(pt[0] * self.width_inch / 1000.0),
                Inches(pt[1] * self.height_inch / 1000.0),
            )
            for pt in points[1:]
        ]
        ff_builder.add_line_segments(emu_points, close=True)
        shape = ff_builder.convert_to_shape()
        if rotation != 0.0:
            shape.rotation = rotation

        if gradient_colors and len(gradient_colors) >= 2:
            apply_gradient_fill(shape._element.spPr, gradient_colors[0], gradient_colors[1], angle=gradient_angle)
        elif bg_color and bg_color.lower() != "transparent":
            shape.fill.solid()
            shape.fill.fore_color.rgb = parse_color(bg_color)
        else:
            shape.fill.background()

        if border_color and border_color.lower() != "transparent":
            shape.line.color.rgb = parse_color(border_color)
            shape.line.width = Pt(border_width_pt)
        else:
            shape.line.fill.background()

        return shape

    def add_badge(
        self,
        box: List[float],
        text: str,
        bg_color: str = "#DBEAFE",
        text_color: str = "#1E40AF",
        font_size: int = 10,
        bold: bool = True,
        border_color: Optional[str] = None,
        border_width_pt: float = 1.0,
        gradient_colors: Optional[List[str]] = None,
        gradient_angle: float = 90.0,
        rotation: float = 0.0,
    ) -> Any:
        """Add a compact status pill badge with solid or gradient fill and optional border."""
        l, t, w, h = self._to_emu(box)
        shape = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
        if rotation != 0.0:
            shape.rotation = rotation
        
        if gradient_colors and len(gradient_colors) >= 2:
            apply_gradient_fill(shape._element.spPr, gradient_colors[0], gradient_colors[1], angle=gradient_angle)
        elif bg_color and bg_color.lower() != "transparent":
            shape.fill.solid()
            shape.fill.fore_color.rgb = parse_color(bg_color)
        else:
            shape.fill.background()

        if border_color and border_color.lower() != "transparent":
            shape.line.color.rgb = parse_color(border_color)
            shape.line.width = Pt(border_width_pt)
        else:
            shape.line.fill.background()

        tf = shape.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = parse_color(text_color)
        return shape

    def add_textbox(
        self,
        box: List[float],
        text: Optional[str] = None,
        runs: Optional[List[Dict[str, Any]]] = None,
        align: str = "left",
        font_size: Union[int, float] = 12,
        font_color: str = "#334155",
        bold: bool = False,
        margin_pt: float = 0.0,
        rotation: float = 0.0,
        word_wrap: bool = True,
        auto_fit_font: bool = False,
        max_lines: Optional[int] = None,
    ) -> Any:
        """Add a text box with structured text runs or simple text with auto-fitting support."""
        l, t, w, h = self._to_emu(box)
        tx_box = self.slide.shapes.add_textbox(l, t, w, h)
        if rotation != 0.0:
            tx_box.rotation = rotation
        tf = tx_box.text_frame
        tf.word_wrap = word_wrap
        margin = Pt(margin_pt)
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = margin

        # Calculate safe auto-fit font size if requested
        target_w_pt = (box[2] * self.width_inch / 1000.0) * 72.0
        effective_font_size = font_size
        if auto_fit_font and text:
            effective_font_size = calculate_safe_font_size(
                text=text,
                target_width_pt=target_w_pt,
                desired_font_size=float(font_size),
                max_lines=max_lines or (1 if not word_wrap else 2),
            )

        align_map = {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
            "justify": PP_ALIGN.JUSTIFY,
        }
        paragraph_align = align_map.get(align.lower(), PP_ALIGN.LEFT)

        if runs:
            p = tf.paragraphs[0]
            p.alignment = paragraph_align
            for run_def in runs:
                run_text = run_def.get("text", "")
                r_size = run_def.get("size", effective_font_size)
                if auto_fit_font and run_text:
                    r_size = calculate_safe_font_size(
                        text=run_text,
                        target_width_pt=target_w_pt,
                        desired_font_size=float(r_size),
                        max_lines=max_lines or 1,
                    )
                if "\n" in run_text and run_def.get("new_paragraph_on_newline", False):
                    lines = run_text.split("\n")
                    for idx, line in enumerate(lines):
                        if idx > 0:
                            p = tf.add_paragraph()
                            p.alignment = paragraph_align
                        r = p.add_run()
                        r.text = line
                        r.font.size = Pt(r_size)
                        r.font.bold = run_def.get("bold", bold)
                        r.font.color.rgb = parse_color(run_def.get("color", font_color))
                else:
                    r = p.add_run()
                    r.text = run_text
                    r.font.size = Pt(r_size)
                    r.font.bold = run_def.get("bold", bold)
                    r.font.color.rgb = parse_color(run_def.get("color", font_color))
        elif text:
            paragraphs = text.split("\n")
            for idx, para_text in enumerate(paragraphs):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.alignment = paragraph_align
                r = p.add_run()
                r.text = para_text
                r.font.size = Pt(effective_font_size)
                r.font.bold = bold
                r.font.color.rgb = parse_color(font_color)

        return tx_box

    def add_bullet_list(
        self,
        box: List[float],
        items: List[Union[str, Dict[str, Any]]],
        font_size: int = 11,
        font_color: str = "#475569",
        bullet_char: str = "•",
        space_after_pt: int = 4,
    ) -> Any:
        """Add clean, consulting-style bullet list items."""
        l, t, w, h = self._to_emu(box)
        tx_box = self.slide.shapes.add_textbox(l, t, w, h)
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        for idx, item in enumerate(items):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.space_after = Pt(space_after_pt)

            if isinstance(item, dict):
                bullet_r = p.add_run()
                bullet_r.text = f"{item.get('bullet', bullet_char)}  "
                bullet_r.font.bold = True
                bullet_r.font.color.rgb = parse_color(item.get("bullet_color", "#2563EB"))

                if "title" in item:
                    title_r = p.add_run()
                    title_r.text = f"{item['title']}: "
                    title_r.font.bold = True
                    title_r.font.size = Pt(item.get("size", font_size))
                    title_r.font.color.rgb = parse_color(item.get("title_color", "#1E293B"))

                text_r = p.add_run()
                text_r.text = str(item.get("text", ""))
                text_r.font.size = Pt(item.get("size", font_size))
                text_r.font.color.rgb = parse_color(item.get("color", font_color))
            else:
                bullet_r = p.add_run()
                bullet_r.text = f"{bullet_char}  "
                bullet_r.font.bold = True
                bullet_r.font.color.rgb = parse_color("#2563EB")

                text_r = p.add_run()
                text_r.text = str(item)
                text_r.font.size = Pt(font_size)
                text_r.font.color.rgb = parse_color(font_color)

        return tx_box

    def add_chart(
        self,
        box: List[float],
        chart_type: str = "column",
        categories: Optional[List[str]] = None,
        series: Optional[List[Dict[str, Any]]] = None,
        title: Optional[str] = None,
        has_legend: bool = True,
        legend_pos: str = "top",
    ) -> Any:
        """Add native PowerPoint editable chart (column, bar, line, pie, doughnut)."""
        chart_type_map = {
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE,
            "line_markers": XL_CHART_TYPE.LINE_MARKERS,
            "pie": XL_CHART_TYPE.PIE,
            "doughnut": XL_CHART_TYPE.DOUGHNUT,
            "area": XL_CHART_TYPE.AREA,
        }
        xl_chart_type = chart_type_map.get(chart_type.lower(), XL_CHART_TYPE.COLUMN_CLUSTERED)

        chart_data = CategoryChartData()
        chart_data.categories = categories or ["Item 1", "Item 2", "Item 3"]

        if series:
            for s in series:
                s_name = s.get("name", "Series")
                s_vals = s.get("values", [10, 20, 30])
                chart_data.add_series(s_name, tuple(s_vals))
        else:
            chart_data.add_series("Series 1", (10, 25, 18))

        l, t, w, h = self._to_emu(box)
        chart_shape = self.slide.shapes.add_chart(xl_chart_type, l, t, w, h, chart_data)
        chart = chart_shape.chart

        chart.has_legend = has_legend
        if has_legend:
            legend_pos_map = {
                "top": XL_LEGEND_POSITION.TOP,
                "bottom": XL_LEGEND_POSITION.BOTTOM,
                "left": XL_LEGEND_POSITION.LEFT,
                "right": XL_LEGEND_POSITION.RIGHT,
            }
            chart.legend.position = legend_pos_map.get(legend_pos.lower(), XL_LEGEND_POSITION.TOP)
            chart.legend.include_in_layout = False

        if title:
            chart.has_title = True
            chart.chart_title.text_frame.text = title

        return chart_shape

    def add_table(
        self,
        box: List[float],
        headers: List[str],
        rows: List[List[Any]],
        col_widths: Optional[List[float]] = None,
        header_bg: str = "#1E293B",
        header_color: str = "#FFFFFF",
        alt_row_bg: Optional[str] = "#F8FAFC",
        row_bg_colors: Optional[List[str]] = None,
        font_size: int = 11,
    ) -> Any:
        """Add native PowerPoint editable table with styled header and custom/zebra row striping."""
        num_rows = len(rows) + 1
        num_cols = len(headers)
        l, t, w, h = self._to_emu(box)

        table_shape = self.slide.shapes.add_table(num_rows, num_cols, l, t, w, h)
        table = table_shape.table

        # Set column widths if provided
        if col_widths and len(col_widths) == num_cols:
            total_w = sum(col_widths)
            table_w_emu = w
            for idx, cw in enumerate(col_widths):
                table.columns[idx].width = int(table_w_emu * (cw / total_w))

        # Format header row
        for col_idx, header_text in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = str(header_text)
            cell.fill.solid()
            cell.fill.fore_color.rgb = parse_color(header_bg)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.bold = True
                    r.font.size = Pt(font_size)
                    r.font.color.rgb = parse_color(header_color)

        # Format data rows
        for row_idx, row_data in enumerate(rows):
            current_row_idx = row_idx + 1
            if row_bg_colors and row_idx < len(row_bg_colors) and row_bg_colors[row_idx]:
                bg = parse_color(row_bg_colors[row_idx])
            else:
                is_alt = (row_idx % 2 == 1)
                bg = parse_color(alt_row_bg) if (is_alt and alt_row_bg) else parse_color("#FFFFFF")

            for col_idx, cell_value in enumerate(row_data):
                if col_idx >= num_cols:
                    break
                cell = table.cell(current_row_idx, col_idx)
                cell.text = str(cell_value)
                if bg:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = bg
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE

                for p in cell.text_frame.paragraphs:
                    # Align numbers to right, text to left
                    val_str = str(cell_value).strip()
                    if val_str.replace(".", "", 1).replace("%", "").replace("$", "").isdigit():
                        p.alignment = PP_ALIGN.RIGHT
                    else:
                        p.alignment = PP_ALIGN.LEFT
                    for r in p.runs:
                        r.font.size = Pt(font_size)
                        r.font.color.rgb = parse_color("#334155")

        return table_shape

    def add_image(self, box: List[float], image_path: str) -> Optional[Any]:
        """Add cropped image, photo, or logo from local file system."""
        if not os.path.exists(image_path):
            print(f"[Warning] Image file not found: {image_path}")
            return None
        l, t, w, h = self._to_emu(box)
        return self.slide.shapes.add_picture(image_path, l, t, w, h)

    def add_progress_bar(
        self,
        box: List[float],
        pct: float,
        bar_color: str = "#C2410C",
        gradient_colors: Optional[List[str]] = None,
        gradient_angle: float = 0.0,
        bg_color: str = "#FED7AA",
        show_text: bool = True,
        text: Optional[str] = None,
        font_size: int = 9,
        text_color: str = "#334155",
        node_glow: bool = True,
    ) -> None:
        """Add a high-fidelity pill progress bar with optional gradient fill, glowing node and percentage label."""
        l, t, w, h = box
        text_width = 40 if show_text else 0
        track_w = max(10, w - text_width)
        bar_h = h

        # 1. Background Track
        track_l, track_t, track_w_emu, track_h_emu = self._to_emu([l, t, track_w, bar_h])
        track_shape = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, track_l, track_t, track_w_emu, track_h_emu)
        track_shape.fill.solid()
        track_shape.fill.fore_color.rgb = parse_color(bg_color)
        track_shape.line.fill.background()

        # 2. Progress Fill
        fill_pct = max(0.05, min(1.0, float(pct)))
        fill_w = track_w * fill_pct
        fill_l, fill_t, fill_w_emu, fill_h_emu = self._to_emu([l, t, fill_w, bar_h])
        fill_shape = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, fill_l, fill_t, fill_w_emu, fill_h_emu)
        
        if gradient_colors and len(gradient_colors) >= 2:
            apply_gradient_fill(fill_shape._element.spPr, gradient_colors[0], gradient_colors[1], angle=gradient_angle)
        else:
            fill_shape.fill.solid()
            fill_shape.fill.fore_color.rgb = parse_color(bar_color)
        fill_shape.line.fill.background()

        # 3. Glowing Pearl Node Dot
        if node_glow and fill_w > 8:
            node_d = bar_h * 0.85
            node_l = l + fill_w - (node_d * 0.85)
            node_t = t + (bar_h - node_d) / 2
            nl, nt, nw, nh = self._to_emu([node_l, node_t, node_d, node_d])
            node_shape = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, nl, nt, nw, nh)
            node_shape.fill.solid()
            node_shape.fill.fore_color.rgb = parse_color("#FFFFFF")
            node_shape.line.color.rgb = parse_color(bar_color)
            node_shape.line.width = Pt(1.0)

        # 4. Text Label
        if show_text:
            text_str = text or f"{int(pct * 100)}%"
            tx_l = l + track_w + 4
            self.add_textbox(
                box=[tx_l, t - 4, text_width, bar_h + 8],
                text=text_str,
                font_size=font_size,
                font_color=text_color,
                bold=True,
                align="left",
            )

    def add_grid(
        self,
        box: List[float],
        cols: int = 3,
        rows: int = 1,
        gap_x: float = 12.0,
        gap_y: float = 12.0,
    ) -> List[Tuple[float, float, float, float]]:
        """Divide a bounding box into an evenly distributed grid of cell boxes in 0-1000 coordinates."""
        l, t, w, h = normalize_box(box)
        cols = max(1, int(cols))
        rows = max(1, int(rows))

        total_gap_w = (cols - 1) * gap_x
        total_gap_h = (rows - 1) * gap_y

        cell_w = (w - total_gap_w) / float(cols)
        cell_h = (h - total_gap_h) / float(rows)

        grid_cells: List[Tuple[float, float, float, float]] = []
        for r in range(rows):
            for c in range(cols):
                cell_left = l + c * (cell_w + gap_x)
                cell_top = t + r * (cell_h + gap_y)
                grid_cells.append((cell_left, cell_top, cell_w, cell_h))

        return grid_cells

    def add_stack(
        self,
        box: List[float],
        direction: str = "vertical",
        gap: float = 8.0,
        align: str = "center",
        children: Optional[List[Dict[str, Any]]] = None,
        bg_color: Optional[str] = None,
        border_color: Optional[str] = None,
        border_width_pt: float = 1.0,
        radius: bool = True,
        gradient_colors: Optional[List[str]] = None,
        gradient_angle: float = 90.0,
    ) -> List[Any]:
        """Add a responsive Flex/Stack container that automatically arranges children with equal spacing and alignment."""
        l, t, w, h = normalize_box(box)
        rendered_shapes: List[Any] = []

        # Optional background container card
        if bg_color or gradient_colors or border_color:
            container = self.add_card(
                box=[l, t, w, h],
                bg_color=bg_color or "transparent",
                border_color=border_color,
                border_width_pt=border_width_pt,
                radius=radius,
                gradient_colors=gradient_colors,
                gradient_angle=gradient_angle,
            )
            rendered_shapes.append(container)

        if not children:
            return rendered_shapes

        # Layout calculations
        is_vertical = direction.lower() == "vertical"

        if is_vertical:
            curr_y = t + 8.0  # Top padding
            for item in children:
                item_type = item.get("type", "text").lower()
                item_h = float(item.get("height", 24.0))
                item_w = float(item.get("width", w - 16.0))

                # Align X
                if align == "center":
                    item_x = l + (w - item_w) / 2.0
                elif align == "right":
                    item_x = l + w - item_w - 8.0
                else:  # left
                    item_x = l + 8.0

                item_box = [item_x, curr_y, item_w, item_h]

                if item_type == "badge":
                    shape = self.add_badge(
                        box=item_box,
                        text=item.get("text", ""),
                        bg_color=item.get("bg_color", "#DBEAFE"),
                        text_color=item.get("text_color", "#1E40AF"),
                        font_size=int(item.get("font_size", 9)),
                        bold=item.get("bold", True),
                    )
                    rendered_shapes.append(shape)
                elif item_type == "card":
                    shape = self.add_card(
                        box=item_box,
                        bg_color=item.get("bg_color", "#FFFFFF"),
                        border_color=item.get("border_color", "#E2E8F0"),
                        radius=item.get("radius", True),
                    )
                    rendered_shapes.append(shape)
                else:  # Default text
                    shape = self.add_textbox(
                        box=item_box,
                        text=item.get("text", ""),
                        font_size=item.get("font_size", 12),
                        font_color=item.get("font_color", "#0F172A"),
                        bold=item.get("bold", False),
                        align=item.get("align", align),
                        word_wrap=item.get("word_wrap", False),
                        auto_fit_font=item.get("auto_fit_font", True),
                    )
                    rendered_shapes.append(shape)

                curr_y += item_h + gap
        else:
            curr_x = l + 8.0  # Left padding
            for item in children:
                item_type = item.get("type", "text").lower()
                item_w = float(item.get("width", 80.0))
                item_h = float(item.get("height", h - 16.0))

                # Align Y
                if align == "middle" or align == "center":
                    item_y = t + (h - item_h) / 2.0
                elif align == "bottom":
                    item_y = t + h - item_h - 8.0
                else:  # top
                    item_y = t + 8.0

                item_box = [curr_x, item_y, item_w, item_h]

                if item_type == "badge":
                    shape = self.add_badge(
                        box=item_box,
                        text=item.get("text", ""),
                        bg_color=item.get("bg_color", "#DBEAFE"),
                        text_color=item.get("text_color", "#1E40AF"),
                        font_size=int(item.get("font_size", 9)),
                        bold=item.get("bold", True),
                    )
                    rendered_shapes.append(shape)
                else:
                    shape = self.add_textbox(
                        box=item_box,
                        text=item.get("text", ""),
                        font_size=item.get("font_size", 12),
                        font_color=item.get("font_color", "#0F172A"),
                        bold=item.get("bold", False),
                        align=item.get("align", "center"),
                        word_wrap=item.get("word_wrap", False),
                        auto_fit_font=item.get("auto_fit_font", True),
                    )
                    rendered_shapes.append(shape)

                curr_x += item_w + gap

        return rendered_shapes

    def add_flex_card(
        self,
        box: List[float],
        title: Optional[str] = None,
        subtitle: Optional[str] = None,
        badge: Optional[str] = None,
        badge_bg: str = "#3B82F6",
        badge_color: str = "#FFFFFF",
        kpi_value: Optional[str] = None,
        kpi_label: Optional[str] = None,
        body_items: Optional[List[str]] = None,
        bg_color: str = "#FFFFFF",
        border_color: str = "#E2E8F0",
        radius: bool = True,
        gradient_colors: Optional[List[str]] = None,
        gradient_angle: float = 90.0,
    ) -> List[Any]:
        """Add a composite semantic card with automatically calculated layout and auto-fitting typography."""
        l, t, w, h = normalize_box(box)
        shapes = []

        # Background Card
        shapes.append(
            self.add_card(
                box=[l, t, w, h],
                bg_color=bg_color,
                border_color=border_color,
                radius=radius,
                gradient_colors=gradient_colors,
                gradient_angle=gradient_angle,
            )
        )

        children_stack: List[Dict[str, Any]] = []

        if badge:
            badge_w = min(w * 0.75, estimate_text_width_pt(badge, 8.5) * 1.5 + 16.0)
            children_stack.append({
                "type": "badge",
                "text": badge,
                "bg_color": badge_bg,
                "text_color": badge_color,
                "font_size": Tokens.FONT_BADGE,
                "width": max(50.0, badge_w),
                "height": 22.0,
            })

        if kpi_value:
            children_stack.append({
                "type": "text",
                "text": kpi_value,
                "font_size": Tokens.FONT_KPI_VAL,
                "bold": True,
                "align": "center",
                "height": 38.0,
                "auto_fit_font": True,
            })

        if kpi_label:
            children_stack.append({
                "type": "text",
                "text": kpi_label,
                "font_size": Tokens.FONT_BODY_SM,
                "font_color": "#64748B",
                "align": "center",
                "height": 18.0,
                "auto_fit_font": True,
            })

        if title:
            children_stack.append({
                "type": "text",
                "text": title,
                "font_size": Tokens.FONT_CARD_TITLE,
                "bold": True,
                "align": "left",
                "height": 24.0,
                "auto_fit_font": True,
            })

        if subtitle:
            children_stack.append({
                "type": "text",
                "text": subtitle,
                "font_size": Tokens.FONT_BODY,
                "font_color": "#64748B",
                "align": "left",
                "height": 20.0,
                "auto_fit_font": True,
            })

        if body_items:
            for item in body_items:
                children_stack.append({
                    "type": "text",
                    "text": f"• {item}",
                    "font_size": Tokens.FONT_BODY_SM,
                    "font_color": "#475569",
                    "align": "left",
                    "height": 18.0,
                    "auto_fit_font": True,
                })

        stack_shapes = self.add_stack(
            box=[l + 6, t + 6, w - 12, h - 12],
            direction="vertical",
            gap=4.0,
            align="center" if kpi_value else "left",
            children=children_stack,
        )
        shapes.extend(stack_shapes)
        return shapes

    def add_kpi_card(
        self,
        box: List[float],
        top_tag: str,
        value: str,
        bottom_badge: str,
        theme_color: str = "#0284C7",
        bg_color: str = "#F0F9FF",
        border_color: str = "#BAE6FD",
        tag_bg: str = "#E0F2FE",
    ) -> None:
        """Add a composite 3-part stacked KPI indicator card (backward compatible)."""
        l, t, w, h = box
        # Container Card
        self.add_card(box=[l, t, w, h], bg_color=bg_color, border_color=border_color, border_width_pt=1.0, radius=True)
        # Top tag
        tag_w = min(w * 0.65, 75)
        tag_h = h * 0.20
        self.add_badge(box=[l + (w - tag_w) / 2, t + 6, tag_w, tag_h], text=top_tag, bg_color=tag_bg, text_color=theme_color, font_size=8)
        # Middle giant value
        val_h = h * 0.45
        self.add_textbox(box=[l + 5, t + tag_h + 4, w - 10, val_h], text=value, font_size=17, font_color="#0F172A", bold=True, align="center", auto_fit_font=True)
        # Bottom pill badge
        badge_w = min(w * 0.75, 85)
        badge_h = h * 0.22
        self.add_badge(box=[l + (w - badge_w) / 2, t + h - badge_h - 5, badge_w, badge_h], text=bottom_badge, bg_color=theme_color, text_color="#FFFFFF", font_size=8, bold=True)

    def save(self, output_path: str) -> str:
        """Save the slide presentation to disk."""
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        self.prs.save(output_path)
        print(f"[Success] Presentation saved to: {output_path}")
        return output_path
