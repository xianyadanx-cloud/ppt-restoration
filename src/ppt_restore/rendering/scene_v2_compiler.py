"""Lower SceneSpec v2 semantic nodes into deterministic PPT operations."""

from __future__ import annotations

import math
import tempfile
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence, Union

from ppt_restore.contracts.schema_v2 import (
    NodeKind,
    RenderOp,
    RenderPlan,
    RenderStrategy,
    SceneNode,
    SceneSpecV2,
    scene_sha256,
)
from ppt_restore.platform.provenance import sha256_file
from ppt_restore.quality.build_manifest import create_build_manifest
from ppt_restore.rendering.compiler import NativePrimitives, set_run_typeface
from ppt_restore.rendering.geometry import box_px_to_emu

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
except ImportError:  # pragma: no cover - compiler tests skip without pptx
    Presentation = None


def _style(node: SceneNode) -> Dict[str, Any]:
    return dict(node.style or {})


def _payload(node: SceneNode) -> Dict[str, Any]:
    return dict(node.payload or {})


def _op(
    node: SceneNode,
    op_type: str,
    payload: Optional[Mapping[str, Any]] = None,
    style: Optional[Mapping[str, Any]] = None,
    suffix: str = "",
) -> RenderOp:
    return RenderOp(
        id=node.id + suffix,
        op_type=op_type,
        bbox_px=node.bbox_px,
        z_index=node.z_index,
        source_node_id=node.id,
        payload=dict(payload or {}),
        style=dict(style or _style(node)),
    )


def _selected_nodes(
    scene: SceneSpecV2, blocks: Optional[Sequence[str]] = None
) -> List[SceneNode]:
    if blocks is None:
        return list(scene.nodes)
    wanted = {str(item) for item in blocks}
    known = {str(item.get("id")) for item in (scene.blocks or ())}
    unknown = sorted(wanted - known) if known else []
    if unknown:
        raise ValueError("unknown SceneSpec block(s): %s" % ", ".join(unknown))
    selected = []
    for node in scene.nodes:
        # Unassigned container/decoration nodes are page-level scaffolding and
        # remain visible in isolated/cumulative previews.
        if node.block_id in wanted or (
            node.block_id is None and node.role in ("container", "decoration")
        ):
            selected.append(node)
    return selected


def lower_scene(
    scene: SceneSpecV2, blocks: Optional[Sequence[str]] = None
) -> RenderPlan:
    """Convert semantic nodes to a stable list of rendering operations."""

    operations: List[RenderOp] = []
    warnings: List[str] = []
    for node in sorted(
        _selected_nodes(scene, blocks), key=lambda item: (int(item.z_index), item.id)
    ):
        kind = str(node.kind)
        payload = _payload(node)
        style = _style(node)
        if kind == NodeKind.GROUP.value:
            continue
        if kind == NodeKind.TEXT.value:
            operations.append(
                _op(
                    node,
                    "text",
                    {"text": payload.get("text", ""), "runs": payload.get("runs", [])},
                    style,
                )
            )
        elif kind == NodeKind.SHAPE.value:
            operations.append(
                _op(node, str(payload.get("shape_type", "rect")), payload, style)
            )
        elif kind == NodeKind.LINE.value:
            operations.append(_op(node, "line", payload, style))
        elif kind == NodeKind.IMAGE.value:
            if str(node.render_strategy) != RenderStrategy.RASTER.value:
                warnings.append(
                    "%s image node requested native rendering; lowered as raster"
                    % node.id
                )
            operations.append(_op(node, "image", payload, style))
        elif kind in (NodeKind.TABLE.value, NodeKind.CHART.value):
            operations.append(_op(node, kind, payload, style))
        elif kind == NodeKind.KPI_CARD.value:
            operations.append(_op(node, "rounded_rect", {}, style, "-base"))
            x, y, width, height = node.bbox_px
            label_h = min(28.0, height * 0.24)
            value_h = min(42.0, height * 0.40)
            operations.append(
                RenderOp(
                    node.id + "-label",
                    "text",
                    (x, y, width, label_h),
                    node.z_index + 1,
                    node.id,
                    {"text": payload.get("label", "")},
                    {
                        "font_size_pt": 11,
                        "weight": "bold",
                        "color": style.get("label_color", "#1B5B9E"),
                        "align": "center",
                        "vertical_anchor": "middle",
                    },
                )
            )
            operations.append(
                RenderOp(
                    node.id + "-value",
                    "text",
                    (x, y + label_h, width, value_h),
                    node.z_index + 1,
                    node.id,
                    {"text": payload.get("value", "")},
                    {
                        "font_size_pt": 22,
                        "weight": "bold",
                        "color": style.get("value_color", "#0F172A"),
                        "align": "center",
                        "vertical_anchor": "middle",
                    },
                )
            )
            if payload.get("delta") is not None:
                operations.append(
                    RenderOp(
                        node.id + "-delta",
                        "rounded_rect",
                        (
                            x + width * 0.12,
                            y + height * 0.73,
                            width * 0.76,
                            height * 0.18,
                        ),
                        node.z_index + 1,
                        node.id,
                        {},
                        {
                            "fill": style.get("delta_fill", "#2563EB"),
                            "corner_radius_px": 8,
                        },
                    )
                )
                operations.append(
                    RenderOp(
                        node.id + "-delta-text",
                        "text",
                        (
                            x + width * 0.12,
                            y + height * 0.73,
                            width * 0.76,
                            height * 0.18,
                        ),
                        node.z_index + 2,
                        node.id,
                        {"text": payload.get("delta", "")},
                        {
                            "font_size_pt": 10,
                            "weight": "bold",
                            "color": "#FFFFFF",
                            "align": "center",
                            "vertical_anchor": "middle",
                        },
                    )
                )
        elif kind == NodeKind.PROGRESS_BAR.value:
            x, y, width, height = node.bbox_px
            value = float(payload.get("value", 0.0))
            minimum = float(payload.get("min", 0.0))
            maximum = float(payload.get("max", 1.0))
            fraction = (
                0.0
                if maximum <= minimum
                else max(0.0, min(1.0, (value - minimum) / (maximum - minimum)))
            )
            track_style = {
                "fill": style.get("background", "#FCE6DC"),
                "corner_radius_px": min(height / 2, 10),
            }
            fill_style = {
                "fill": style.get("fill", "#C2410C"),
                "corner_radius_px": min(height / 2, 10),
            }
            if style.get("track_gradient_stops"):
                track_style["gradient_stops"] = style["track_gradient_stops"]
                track_style["gradient_angle"] = style.get("gradient_angle", 0)
            if style.get("gradient_stops"):
                fill_style["gradient_stops"] = style["gradient_stops"]
                fill_style["gradient_angle"] = style.get("gradient_angle", 0)
            operations.append(_op(node, "rounded_rect", {}, track_style, "-track"))
            fill_width = max(1.0, width * fraction)
            operations.append(
                RenderOp(
                    node.id + "-fill",
                    "rounded_rect",
                    (x, y, fill_width, height),
                    node.z_index + 1,
                    node.id,
                    {},
                    fill_style,
                )
            )
            marker = payload.get("marker", style.get("marker", False))
            if marker:
                marker_size = float(style.get("marker_size_px", max(height * 0.9, 6.0)))
                marker_style = {
                    "fill": style.get("marker_fill", "#FFFFFF"),
                    "stroke": style.get("marker_stroke", style.get("fill", "#C2410C")),
                    "stroke_width_px": style.get("marker_stroke_width_px", 1.0),
                }
                if style.get("marker_shadow"):
                    marker_style["shadow"] = style["marker_shadow"]
                operations.append(
                    RenderOp(
                        node.id + "-marker",
                        "circle",
                        (
                            x + max(0.0, fill_width - marker_size / 2.0),
                            y + (height - marker_size) / 2.0,
                            marker_size,
                            marker_size,
                        ),
                        node.z_index + 2,
                        node.id,
                        {},
                        marker_style,
                    )
                )
            if payload.get("label") is not None:
                operations.append(
                    RenderOp(
                        node.id + "-label",
                        "text",
                        node.bbox_px,
                        node.z_index + 2,
                        node.id,
                        {"text": payload.get("label", "")},
                        {
                            "font_size_pt": style.get("font_size_pt", 10),
                            "weight": "bold",
                            "color": style.get("label_color", "#9A3412"),
                            "align": "center",
                            "vertical_anchor": "middle",
                        },
                    )
                )
        elif kind == NodeKind.BADGE.value:
            operations.append(_op(node, "rounded_rect", {}, style, "-base"))
            operations.append(
                RenderOp(
                    node.id + "-text",
                    "text",
                    node.bbox_px,
                    node.z_index + 1,
                    node.id,
                    {"text": payload.get("text", "")},
                    {
                        "font_size_pt": style.get("font_size_pt", 10),
                        "weight": "bold",
                        "color": style.get("text_color", "#000000"),
                        "align": "center",
                        "vertical_anchor": "middle",
                    },
                )
            )
        elif kind == NodeKind.PROCESS.value:
            steps = payload.get("steps", [])
            if not isinstance(steps, list) or not steps:
                warnings.append("%s process has no steps" % node.id)
            else:
                x, y, width, height = node.bbox_px
                step_width = width / float(len(steps))
                for index, step in enumerate(steps):
                    text = (
                        step.get("text", "") if isinstance(step, Mapping) else str(step)
                    )
                    operations.append(
                        RenderOp(
                            "%s-step-%d" % (node.id, index + 1),
                            "text",
                            (x + index * step_width, y, step_width, height),
                            node.z_index + index,
                            node.id,
                            {"text": text},
                            {
                                "font_size_pt": style.get("font_size_pt", 12),
                                "color": style.get("color", "#0F172A"),
                                "align": "center",
                                "vertical_anchor": "middle",
                            },
                        )
                    )
        else:
            warnings.append("unsupported node kind %s" % kind)
    # Composite text inherits the requested font rather than silently using Arial.
    from dataclasses import replace

    source_nodes = {node.id: node for node in scene.nodes}
    inherited = []
    for operation in operations:
        owner = source_nodes.get(operation.source_node_id)
        if owner and operation.op_type == "text":
            merged = {
                key: value
                for key, value in (owner.style or {}).items()
                if key
                in (
                    "font_family",
                    "line_spacing",
                    "line_spacing_pt",
                    "space_before_pt",
                    "space_after_pt",
                    "word_wrap",
                )
            }
            merged.update(operation.style)
            operation = replace(operation, style=merged)
        inherited.append(operation)
    return RenderPlan(
        "1.0", scene_sha256(scene), scene.canvas, tuple(inherited), tuple(warnings)
    )


class SceneSpecCompiler(NativePrimitives):
    """Compile validated SceneSpec v2 using native drawing primitives."""

    VERSION = "scene-v2-1"

    def compile(
        self,
        scene: SceneSpecV2,
        output_path: Optional[Union[str, Path]] = None,
        *,
        canonical_path: Optional[Union[str, Path]] = None,
        case_dir: Optional[Union[str, Path]] = None,
        render_strategy: str = "hybrid_editable",
        strict_native: bool = False,
        blocks: Optional[Sequence[str]] = None,
    ):
        if Presentation is None:
            raise RuntimeError("python-pptx is required to compile a SceneSpecV2")
        if not isinstance(scene, SceneSpecV2):
            raise TypeError("SceneSpecCompiler requires SceneSpecV2")
        if scene.state != "VALIDATED":
            raise ValueError(
                "SceneSpec must be VALIDATED before build; state=%s" % scene.state
            )
        if canonical_path is not None:
            canonical = Path(canonical_path)
            if not canonical.is_file():
                raise FileNotFoundError(str(canonical))
            if scene.canonical_sha256 != sha256_file(canonical):
                raise ValueError(
                    "canonical_sha256 mismatch; refusing to build stale SceneSpec"
                )
        if case_dir is not None:
            evidence_path = Path(case_dir) / "evidence.json"
            if evidence_path.is_file() and scene.evidence_sha256 != sha256_file(
                evidence_path
            ):
                raise ValueError(
                    "evidence_sha256 mismatch; refusing to build stale SceneSpec"
                )
        plan = lower_scene(scene, blocks=blocks)
        if strict_native or render_strategy == "strict_native":
            if any(op.op_type == "image" for op in plan.operations):
                raise ValueError("strict_native cannot compile raster image operations")
        canvas = self._scene_canvas(scene)
        self.canvas = canvas
        self._warnings = list(plan.warnings)
        prs = Presentation()
        prs.slide_width, prs.slide_height = canvas.width_emu, canvas.height_emu
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        temp_root = tempfile.TemporaryDirectory(prefix="pptrestore-v2-")
        try:
            image_path = Path(canonical_path) if canonical_path else None
            count = 0
            for op in sorted(
                plan.operations, key=lambda item: (int(item.z_index), item.id)
            ):
                try:
                    if op.op_type == "image":
                        if image_path is None or not image_path.is_file():
                            raise ValueError("image operation requires canonical_path")
                        self._compile_image(slide, op, image_path, Path(temp_root.name))
                    elif op.op_type == "table":
                        self._compile_table(slide, op, canvas)
                    elif op.op_type == "chart":
                        self._compile_chart(slide, op, canvas)
                    else:
                        self._compile_render_op(slide, op, canvas)
                    count += 1
                except Exception as exc:
                    self._warnings.append("%s: %s" % (op.id, exc))
            result = _build_result(self, output_path, prs, scene, count, canvas)
        finally:
            temp_root.cleanup()
        if (
            output_path is not None
            and case_dir is not None
            and canonical_path is not None
        ):
            plan_path = Path(case_dir) / "render_plan.json"
            plan_path.parent.mkdir(parents=True, exist_ok=True)
            plan_path.write_text(plan.to_json(indent=2) + "\n", encoding="utf-8")
            result.render_plan_path = str(plan_path)
            evidence_path = Path(case_dir) / "evidence.json"
            scene_path = Path(case_dir) / "scene.json"
            if (
                evidence_path.is_file()
                and scene_path.is_file()
                and Path(output_path).is_file()
            ):
                manifest, manifest_path = create_build_manifest(
                    case_dir,
                    output_path,
                    canonical_path=canonical_path,
                    evidence_path=evidence_path,
                    scene_path=scene_path,
                    render_plan_path=plan_path,
                    compiler_version=self.VERSION,
                )
                result.manifest_path = str(manifest_path)
        return result

    def _compile_render_op(self, slide, op: RenderOp, canvas) -> None:
        payload = dict(op.payload or {})
        style = dict(op.style or {})
        typ = op.op_type
        if typ == "rect":
            element_type = "rect"
        elif typ in ("rounded_rect", "rounded_rectangle"):
            element_type = "rounded_rect"
        elif typ in ("ellipse", "oval", "circle"):
            element_type = "circle"
        elif typ in ("line", "connector"):
            element_type = "line"
        elif typ == "text":
            element_type = "text"
        else:
            element_type = (
                "freeform" if typ in ("freeform", "polygon", "path") else "rect"
            )
        element = {
            "id": op.id,
            "element_type": element_type,
            "layout_bbox_px": list(op.bbox_px),
            "z_index": op.z_index,
            "shape_style": style,
            "text_style": style,
            "text": payload.get("text"),
            "runs": payload.get("runs", ()),
            "align": style.get("align", "left"),
        }
        points = payload.get("points", payload.get("vertices"))
        if points is None and payload.get("points_norm") is not None:
            # Public SceneSpec geometry is normalized; native primitives use
            # freeform API expects points relative to the node box in pixels.
            x, y, _, _ = op.bbox_px
            points = []
            for point in payload.get("points_norm") or ():
                if not isinstance(point, (list, tuple)) or len(point) != 2:
                    raise ValueError("points_norm must contain [x, y] pairs")
                points.append(
                    (
                        float(point[0]) * canvas.width_px / 1000.0 - x,
                        float(point[1]) * canvas.height_px / 1000.0 - y,
                    )
                )
        if points is not None:
            element["points"] = points
        if element_type == "text" and element["text"] is None:
            element["text"] = ""
        self._compile_element(slide, element, canvas)

    def _compile_image(
        self, slide, op: RenderOp, image_path: Path, temp_root: Path
    ) -> None:
        from PIL import Image

        crop_box = op.payload.get("source_bbox_px", op.bbox_px)
        x, y, width, height = [int(round(value)) for value in crop_box]
        with Image.open(image_path) as image:
            if "source_bbox_norm" in op.payload and "source_bbox_px" not in op.payload:
                norm = op.payload["source_bbox_norm"]
                if (
                    not isinstance(norm, (list, tuple))
                    or len(norm) != 4
                    or any(float(v) < 0 or float(v) > 1000 for v in norm)
                ):
                    raise ValueError("source_bbox_norm must be a 0-1000 box")
                x = int(round(float(norm[0]) * image.width / 1000.0))
                y = int(round(float(norm[1]) * image.height / 1000.0))
                width = int(round(float(norm[2]) * image.width / 1000.0))
                height = int(round(float(norm[3]) * image.height / 1000.0))
            if width <= 0 or height <= 0:
                raise ValueError("image crop width and height must be positive")
            x = max(0, min(image.width - 1, x))
            y = max(0, min(image.height - 1, y))
            width = min(width, image.width - x)
            height = min(height, image.height - y)
            crop = image.convert("RGB").crop((x, y, x + width, y + height))
            path = temp_root / (op.id.replace("/", "_") + ".png")
            crop.save(path)
        left, top, w, h = box_px_to_emu(op.bbox_px, self.canvas)
        slide.shapes.add_picture(str(path), left, top, w, h)

    def _compile_table(self, slide, op: RenderOp, canvas) -> None:
        cells = op.payload.get("cells", [])
        if not isinstance(cells, list) or not cells or not isinstance(cells[0], list):
            raise ValueError("table cells must be a non-empty 2D list")
        rows, cols = len(cells), max(len(row) for row in cells)
        x, y, w, h = box_px_to_emu(op.bbox_px, canvas)
        table = slide.shapes.add_table(rows, cols, x, y, w, h).table

        def dimensions(key, count, total):
            values = op.payload.get(key, [1] * count)
            if len(values) != count or any(
                not math.isfinite(float(v)) or float(v) <= 0 for v in values
            ):
                raise ValueError(
                    "%s must have one positive weight per row/column" % key
                )
            sizes = [round(total * float(v) / sum(map(float, values))) for v in values]
            sizes[-1] += total - sum(sizes)
            return sizes

        for column, width in zip(table.columns, dimensions("column_widths", cols, w)):
            column.width = width
        for row, height in zip(table.rows, dimensions("row_heights", rows, h)):
            row.height = height
        for r in range(rows):
            for c in range(cols):
                value = cells[r][c] if c < len(cells[r]) else ""
                cell = table.cell(r, c)
                cell.text = str(
                    value.get("text", "") if isinstance(value, Mapping) else value
                )
                cell.margin_left = cell.margin_right = cell.margin_top = (
                    cell.margin_bottom
                ) = 0
                cell.vertical_anchor = {
                    "middle": MSO_ANCHOR.MIDDLE,
                    "bottom": MSO_ANCHOR.BOTTOM,
                }.get(op.style.get("vertical_anchor"), MSO_ANCHOR.TOP)
                cell.text_frame.word_wrap = bool(op.style.get("word_wrap", True))
                for paragraph in cell.text_frame.paragraphs:
                    from ppt_restore.rendering.compiler import set_paragraph_spacing

                    set_paragraph_spacing(paragraph, op.style)
                    paragraph.alignment = PP_ALIGN.CENTER
                    for run in paragraph.runs:
                        run.font.size = __import__("pptx").util.Pt(
                            float(op.style.get("font_size_pt", 11))
                        )
                        set_run_typeface(run, str(op.style.get("font_family", "Arial")))
                        run.font.color.rgb = self._rgb(
                            op.style.get(
                                "header_color" if r == 0 else "color", "#000000"
                            )
                        )
                        if op.style.get(
                            "header_weight" if r == 0 else "weight",
                            op.style.get("weight"),
                        ) in ("bold", "700", "800", "900"):
                            run.font.bold = True
                fill = None
                row_colors = op.payload.get(
                    "row_bg_colors", op.style.get("row_bg_colors", [])
                )
                header_fill = op.payload.get("header_bg", op.style.get("header_bg"))
                if r == 0 and header_fill:
                    fill = header_fill
                elif isinstance(row_colors, list) and row_colors:
                    fill = (
                        row_colors[(r - 1 if header_fill else r) % len(row_colors)]
                        if (r - 1 if header_fill else r) >= 0
                        else None
                    )
                if fill is None:
                    fill = op.style.get("fill")
                if fill:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = self._rgb(fill)
        for bounds in op.payload.get("merges", []):
            if len(bounds) != 4 or any(type(v) is not int for v in bounds):
                raise ValueError(
                    "table merge requires [first_row, first_column, last_row, last_column]"
                )
            r1, c1, r2, c2 = bounds
            if not (0 <= r1 <= r2 < rows and 0 <= c1 <= c2 < cols):
                raise ValueError("table merge is outside the table")
            table.cell(r1, c1).merge(table.cell(r2, c2))
        self._table_progress(slide, table, op, x, y)

    def _table_progress(self, slide, table, op, x, y):
        """Recompute native overlays from cell geometry on every build."""
        for index, spec in enumerate(op.payload.get("progress_bars", [])):
            row, column = spec.get("row"), spec.get("column")
            if (
                type(row) is not int
                or type(column) is not int
                or not (0 <= row < len(table.rows) and 0 <= column < len(table.columns))
            ):
                raise ValueError(
                    "table progress row/column must identify an existing cell"
                )
            value = float(spec.get("value", 0))
            padding = float(spec.get("padding_fraction", 0.1))
            thickness = float(spec.get("height_fraction", 0.18))
            center = float(spec.get("center_y_fraction", 0.5))
            if not all(
                math.isfinite(v) for v in (value, padding, thickness, center)
            ) or not (
                0 <= value <= 1
                and 0 <= padding < 0.5
                and 0 < thickness <= 1
                and thickness / 2 <= center <= 1 - thickness / 2
            ):
                raise ValueError(
                    "invalid table progress geometry or value (expected fraction 0..1)"
                )
            cw, rh = table.columns[column].width, table.rows[row].height
            left = (
                x
                + sum(table.columns[i].width for i in range(column))
                + round(cw * padding)
            )
            top = (
                y
                + sum(table.rows[i].height for i in range(row))
                + round(rh * (center - thickness / 2))
            )
            width, height = round(cw * (1 - 2 * padding)), round(rh * thickness)

            def shape(part, kind, sx, sy, sw, sh, color):
                item = slide.shapes.add_shape(kind, sx, sy, max(1, sw), max(1, sh))
                item.name = "%s:progress:%d:%s" % (op.source_node_id, index, part)
                item.fill.solid()
                item.fill.fore_color.rgb = self._rgb(color)
                item.line.fill.background()
                return item

            kind = (
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
                if spec.get("rounded", False)
                else MSO_AUTO_SHAPE_TYPE.RECTANGLE
            )
            track = shape(
                "track",
                kind,
                left,
                top,
                width,
                height,
                spec.get("track_color", "#DFE7F2"),
            )
            if spec.get("rounded", False):
                track.adjustments[0] = 0.5
            if spec.get("track_gradient_stops"):
                self._gradient(
                    track,
                    spec["track_gradient_stops"],
                    float(spec.get("gradient_angle", 0)),
                )
            if value > 0:
                fill = shape(
                    "fill",
                    kind,
                    left,
                    top,
                    round(width * value),
                    height,
                    spec.get("fill_color", "#2563EB"),
                )
                if spec.get("rounded", False):
                    fill.adjustments[0] = 0.5
                if spec.get("gradient_stops"):
                    self._gradient(
                        fill,
                        spec["gradient_stops"],
                        float(spec.get("gradient_angle", 0)),
                    )
            if spec.get("marker", False):
                # EMU dimensions are physical, so equal dimensions remain circular.
                fraction = float(spec.get("marker_diameter_fraction", thickness * 1.2))
                if not math.isfinite(fraction) or not 0 < fraction <= 1:
                    raise ValueError("marker_diameter_fraction must be in (0, 1]")
                diameter = min(
                    round(rh * fraction), round(2 * rh * min(center, 1 - center)), cw
                )
                if diameter > 0:
                    cell_left = left - round(cw * padding)
                    marker_left = max(
                        cell_left,
                        min(
                            cell_left + cw - diameter,
                            left + round(width * value) - diameter // 2,
                        ),
                    )
                    shape(
                        "marker",
                        MSO_AUTO_SHAPE_TYPE.OVAL,
                        marker_left,
                        top + height // 2 - diameter // 2,
                        diameter,
                        diameter,
                        spec.get("marker_color", "#FFFFFF"),
                    )

    def _compile_chart(self, slide, op: RenderOp, canvas) -> None:
        payload = op.payload
        chart_data = CategoryChartData()
        chart_data.categories = [str(item) for item in payload.get("categories", [])]
        series = payload.get("series", [])
        for item in series:
            if isinstance(item, Mapping):
                chart_data.add_series(
                    str(item.get("name", "Series")),
                    [float(value) for value in item.get("values", [])],
                )
        chart_type = str(payload.get("chart_type", "column")).casefold()
        kind = (
            XL_CHART_TYPE.BAR_CLUSTERED
            if chart_type in ("bar", "horizontal_bar")
            else XL_CHART_TYPE.LINE
            if chart_type == "line"
            else XL_CHART_TYPE.PIE
            if chart_type == "pie"
            else XL_CHART_TYPE.COLUMN_CLUSTERED
        )
        x, y, w, h = box_px_to_emu(op.bbox_px, canvas)
        chart = slide.shapes.add_chart(kind, x, y, w, h, chart_data).chart
        chart.has_legend = bool(payload.get("has_legend", False))
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM

    @staticmethod
    def _rgb(value):
        text = str(value).lstrip("#")
        if len(text) != 6:
            text = "000000"
        return RGBColor(int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16))


def _build_result(self, output_path, prs, scene, count, canvas):
    """Create the semantic compiler build result."""
    from ppt_restore.rendering.compiler import BuildResult

    result = BuildResult(
        status="BUILT",
        element_count=count,
        slide_width_emu=canvas.width_emu,
        slide_height_emu=canvas.height_emu,
    )
    result.scene_sha256 = scene_sha256(scene)
    result.warnings = list(getattr(self, "_warnings", []))
    if output_path is not None:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(path))
        result.output_path = str(path)
    from ppt_restore.rendering.compiler import _is_picture_shape

    result.picture_count = sum(
        1 for slide in prs.slides for shape in slide.shapes if _is_picture_shape(shape)
    )
    return result


__all__ = ["lower_scene", "SceneSpecCompiler"]
