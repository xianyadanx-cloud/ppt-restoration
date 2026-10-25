"""Native python-pptx drawing primitives for the semantic compiler.

No image shape is ever inserted.  All supported primitives become native
AutoShapes, connectors, freeform geometry, or editable text boxes.
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from typing import Any, List, Mapping, Optional, Sequence

from ppt_restore.rendering.fonts import FontRegistry
from ppt_restore.rendering.geometry import CanvasSpec, box_px_to_emu, px_to_emu

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls, qn
    from pptx.oxml.xmlchemy import OxmlElement
except ImportError:  # allows schema-only imports in minimal environments
    Presentation = None
    MSO_AUTO_SIZE = None
    qn = None


def set_run_typeface(run, family: str) -> None:
    """Set both Latin and East-Asian DrawingML faces on a text run.

    ``python-pptx`` only writes the Latin face for ``font.name``.  WPS and
    PowerPoint may otherwise substitute Chinese glyphs even when the requested
    family is installed, changing line breaks and visual weight.
    """
    family = str(family)
    run.font.name = family
    rpr = run._r.get_or_add_rPr()
    rpr.attrib.pop(qn("a:lang"), None)
    rpr.set("lang", "zh-CN")
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = OxmlElement("a:ea")
        rpr.append(ea)
    ea.set("typeface", family)


def set_paragraph_spacing(paragraph, style):
    """line_spacing is a multiple; line_spacing_pt is explicit point spacing."""
    from pptx.util import Pt

    for key in ("line_spacing", "line_spacing_pt", "space_before_pt", "space_after_pt"):
        value = style.get(key)
        if value is None:
            continue
        value = float(value)
        if (
            not math.isfinite(value)
            or value < 0
            or (key.startswith("line_spacing") and value == 0)
        ):
            raise ValueError(
                "%s must be finite and %s"
                % (key, "positive" if key.startswith("line_spacing") else "nonnegative")
            )
        if key == "line_spacing":
            paragraph.line_spacing = value
        elif key == "line_spacing_pt":
            paragraph.line_spacing = Pt(value)
        elif key == "space_before_pt":
            paragraph.space_before = Pt(value)
        else:
            paragraph.space_after = Pt(value)


@dataclass
class BuildResult:
    output_path: Optional[str] = None
    status: str = "BUILT"
    element_count: int = 0
    picture_count: int = 0
    warnings: List[str] = field(default_factory=list)
    scene_sha256: str = ""
    slide_width_emu: int = 0
    slide_height_emu: int = 0
    # v2 builds bind every intermediate artifact by content hash.  These are
    # optional when compiling a synthetic scene without a case directory.
    manifest_path: Optional[str] = None
    render_plan_path: Optional[str] = None
    content_inventory_path: Optional[str] = None
    content_dom_check_path: Optional[str] = None

    @property
    def path(self):
        return self.output_path

    def to_dict(self):
        data = {
            "output_path": self.output_path,
            "status": self.status,
            "element_count": self.element_count,
            "picture_count": self.picture_count,
            "warnings": list(self.warnings),
            "scene_sha256": self.scene_sha256,
            "slide_width_emu": self.slide_width_emu,
            "slide_height_emu": self.slide_height_emu,
        }
        if self.manifest_path is not None:
            data["manifest_path"] = self.manifest_path
        if self.render_plan_path is not None:
            data["render_plan_path"] = self.render_plan_path
        if self.content_inventory_path is not None:
            data["content_inventory_path"] = self.content_inventory_path
        if self.content_dom_check_path is not None:
            data["content_dom_check_path"] = self.content_dom_check_path
        return data


def _get(obj: Any, key: str, default: Any = None) -> Any:
    if isinstance(obj, Mapping):
        return obj.get(key, default)
    return getattr(obj, key, default)


def _split_alpha(value: Any):
    """Return (value without alpha, alpha 0..1 or None) for #RRGGBBAA / rgba() / RGBA tuples."""
    if isinstance(value, (tuple, list)) and len(value) >= 4:
        a = float(value[3])
        return tuple(value[:3]), (a / 255.0 if a > 1 else a)
    if isinstance(value, str):
        text = value.strip()
        low = text.casefold().replace(" ", "")
        if low.startswith("rgba(") and low.endswith(")"):
            parts = low[5:-1].split(",")
            if len(parts) == 4:
                r, g, b = (int(float(x)) for x in parts[:3])
                a = float(parts[3].rstrip("%")) / (100.0 if parts[3].endswith("%") else 1.0)
                return "%02X%02X%02X" % (r, g, b), max(0.0, min(1.0, a))
        if low.startswith("rgb(") and low.endswith(")"):
            parts = low[4:-1].split(",")
            if len(parts) == 3:
                return "%02X%02X%02X" % tuple(int(float(x)) for x in parts), None
        bare = text.lstrip("#")
        if len(bare) == 8 and all(c in "0123456789abcdefABCDEF" for c in bare):
            return bare[:6], int(bare[6:], 16) / 255.0
    return value, None


def _hex(value: Any, default: str = "000000") -> str:
    value, _ = _split_alpha(value)
    if value is None:
        return default
    if isinstance(value, RGBColor):
        return "%02X%02X%02X" % tuple(value)
    if isinstance(value, (tuple, list)) and len(value) >= 3:
        return "%02X%02X%02X" % (int(value[0]), int(value[1]), int(value[2]))
    value = str(value).strip().lstrip("#")
    names = {
        "white": "FFFFFF",
        "black": "000000",
        "transparent": "000000",
        "blue": "2563EB",
        "gray": "64748B",
    }
    value = names.get(value.casefold(), value)
    return value.upper() if len(value) == 6 else default


def _rgb(value: Any, default: str = "000000") -> RGBColor:
    h = _hex(value, default)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _style(element: Any, name: str, default=None):
    value = _get(element, name, None)
    if value is not None:
        return value
    for key in ("shape_style", "shapeStyle", "text_style", "textStyle", "style"):
        s = _get(element, key, None)
        if s is not None:
            value = _get(s, name, None)
            if value is not None:
                return value
    return default


def strip_theme_style(shape) -> None:
    """Remove the theme ``p:style`` that paints a default outer shadow.

    python-pptx stamps ``effectRef idx="2"`` onto autoshapes.  LibreOffice and
    PowerPoint then draw the theme shadow even when the scene requested none.
    Explicit ``shadow`` / ``glow`` effects live in ``spPr/effectLst`` and stay.
    """

    element = getattr(shape, "_element", None)
    if element is None or qn is None:
        return
    style = element.find(qn("p:style"))
    if style is not None:
        element.remove(style)


# Autoshapes that may carry a text body.  A chevron with a label must stay a
# chevron; the generic "has text → textbox" rule would otherwise flatten it.
_SHAPE_TYPES = {
    "rect": "RECTANGLE",
    "rectangle": "RECTANGLE",
    "rounded_rect": "ROUNDED_RECTANGLE",
    "rounded_rectangle": "ROUNDED_RECTANGLE",
    "rounded_rect_px": "ROUNDED_RECTANGLE",
    "circle": "OVAL",
    "ellipse": "OVAL",
    "oval": "OVAL",
    "circle_px": "OVAL",
    "chevron": "CHEVRON",
    "home_plate": "HOME_PLATE",
    "pentagon": "PENTAGON",
    "right_arrow": "RIGHT_ARROW",
    "left_arrow": "LEFT_ARROW",
    "down_arrow": "DOWN_ARROW",
    "up_arrow": "UP_ARROW",
}


def _bbox(element: Any) -> Sequence[float]:
    for key in ("layout_bbox_px", "bbox_px", "bbox", "box"):
        value = _get(element, key, None)
        if value is not None:
            if _get(element, "coordinate_space", "px") in (
                "norm",
                "normalized",
                "0-1000",
            ):
                raise ValueError(
                    "normalized coordinates must be converted to pixels before compilation"
                )
            return value
    raise ValueError("element is missing layout_bbox_px")


class NativePrimitives:
    def __init__(
        self,
        font_registry: Optional[FontRegistry] = None,
        canvas: Optional[CanvasSpec] = None,
    ):
        self.font_registry = font_registry or FontRegistry(auto_scan=False)
        self.canvas = canvas or CanvasSpec()
        self._warnings: List[str] = []

    def _scene_canvas(self, scene):
        raw = scene.canvas
        wp, hp = int(raw.get("width_px", 1440)), int(raw.get("height_px", 810))
        wi, hi = raw.get("width_in"), raw.get("height_in")
        if wi is None or hi is None:
            wi = 10.0 if abs(wp / hp - 4 / 3) < 0.01 else 13.333333333333334
            hi = wi * hp / wp
        return CanvasSpec(wp, hp, float(wi), float(hi))

    def _compile_element(self, slide, element, canvas):
        typ = (
            str(_get(element, "element_type", _get(element, "type", "rect")))
            .casefold()
            .replace("-", "_")
        )
        bbox = list(_bbox(element))
        if _style(element, "box_is_ink", False) and _get(element, "text", None):
            from ppt_restore.rendering.text_fit import expand_ink_frame

            bbox = list(
                expand_ink_frame(
                    bbox,
                    float(_style(element, "font_size_pt", 12) or 12),
                    canvas,
                    ascent_gap_em=_style(element, "ascent_gap_em", None),
                    descent_pad_em=_style(element, "descent_pad_em", None),
                )
            )
        x, y, w, h = box_px_to_emu(bbox, canvas)
        shaped = typ in _SHAPE_TYPES or typ in (
            "line",
            "connector",
            "freeform",
            "polygon",
            "path",
        )
        if (
            typ in ("text", "textbox", "text_box", "text_box_px")
            or _get(element, "text", None) is not None
            and not shaped
        ):
            shape = slide.shapes.add_textbox(x, y, w, h)
            self._format_text(shape, element, canvas)
            strip_theme_style(shape)
            return shape
        if typ in ("line", "connector"):
            shape = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, x, y, x + w, y + h
            )
            self._format_line(shape, element, canvas)
            strip_theme_style(shape)
            return shape
        if typ in ("freeform", "polygon", "path"):
            pts = _get(element, "points", None) or _get(element, "vertices", None)
            if not pts:
                pts = [(0, 0), (w, 0), (w, h), (0, h)]
            builder = slide.shapes.build_freeform(x, y, scale=(1, 1))
            first = pts[0]
            builder.move_to(
                x + px_to_emu(first[0], "x", canvas),
                y + px_to_emu(first[1], "y", canvas),
            )
            builder.add_line_segments(
                [
                    (x + px_to_emu(p[0], "x", canvas), y + px_to_emu(p[1], "y", canvas))
                    for p in pts[1:]
                ],
                close=True,
            )
            shape = builder.convert_to_shape()
            self._format_shape(shape, element, canvas)
            strip_theme_style(shape)
            return shape
        kind_name = _SHAPE_TYPES.get(typ, "RECTANGLE")
        kind = getattr(MSO_AUTO_SHAPE_TYPE, kind_name, None)
        prst_override = None
        if kind is None and kind_name == "HOME_PLATE":
            # python-pptx omits homePlate.  Pentagon is the flat-left arrow;
            # renaming the preset keeps the geometry on LibreOffice/PowerPoint.
            kind = MSO_AUTO_SHAPE_TYPE.PENTAGON
            prst_override = "homePlate"
        if kind is None:
            kind = MSO_AUTO_SHAPE_TYPE.RECTANGLE
        shape = slide.shapes.add_shape(kind, x, y, w, h)
        if prst_override and qn is not None:
            prst = shape._element.spPr.find(qn("a:prstGeom"))
            if prst is not None:
                prst.set("prst", prst_override)
        self._format_shape(shape, element, canvas)
        self._apply_shape_adjust(shape, element)
        if kind == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE and prst_override is None:
            self._set_corner_adjustment(
                shape, float(_style(element, "corner_radius_px", 0) or 0), w, h, canvas
            )
        if _get(element, "text", None) is not None:
            self._format_text(shape, element, canvas)
        strip_theme_style(shape)
        return shape

    def _format_shape(self, shape, element, canvas):
        fill = _style(element, "fill", _style(element, "fill_color", None))
        opacity = _style(element, "fill_opacity", _style(element, "opacity", None))
        opacity = 1.0 if opacity is None else float(opacity)
        fill_alpha = _split_alpha(fill)[1]
        if fill_alpha is not None:
            opacity *= fill_alpha
        if fill is None or str(fill).casefold() in ("none", "transparent"):
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(fill)
            self._set_alpha(shape.fill._xPr, opacity)
        line_color = _style(element, "line_color", _style(element, "stroke", None))
        line_width = float(
            _style(element, "line_width_px", _style(element, "stroke_width_px", 0)) or 0
        )
        if line_color is None or line_width <= 0:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = _rgb(line_color)
            shape.line.width = px_to_emu(line_width, "x", canvas)
            self._set_alpha(
                shape.line._get_or_add_ln(),
                float(_style(element, "line_opacity", 1.0) or 1.0),
            )
        stops = _style(element, "gradient_stops", None)
        if stops:
            self._gradient(
                shape, stops, float(_style(element, "gradient_angle", 0) or 0)
            )
        self._effects(shape, element, canvas)

    def _format_line(self, shape, element, canvas):
        color = _style(element, "line_color", _style(element, "stroke", "000000"))
        shape.line.color.rgb = _rgb(color)
        width = float(
            _style(element, "line_width_px", _style(element, "stroke_width_px", 1)) or 1
        )
        shape.line.width = px_to_emu(width, "x", canvas)

    def _format_text(self, shape, element, canvas):
        tf = shape.text_frame
        tf.clear()
        # ``add_textbox`` writes ``spAutoFit``.  LibreOffice still wraps on the
        # box width under wrap="none"+spAutoFit, which splits CJK titles.
        if MSO_AUTO_SIZE is not None:
            tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.word_wrap = bool(_style(element, "word_wrap", False))
        margins = (
            _style(element, "margins_pt", None)
            or _style(element, "margin_pt", None)
            or (0, 0, 0, 0)
        )
        if isinstance(margins, (int, float)):
            margins = (margins,) * 4
        if len(margins) == 4:
            tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [
                int(round(float(v) * 12700)) for v in margins
            ]
        anchor = str(
            _style(element, "vertical_anchor", _style(element, "anchor", "top"))
        ).casefold()
        tf.vertical_anchor = {
            "middle": MSO_ANCHOR.MIDDLE,
            "center": MSO_ANCHOR.MIDDLE,
            "bottom": MSO_ANCHOR.BOTTOM,
        }.get(anchor, MSO_ANCHOR.TOP)
        p = tf.paragraphs[0]
        text = _get(element, "text", "") or ""
        set_paragraph_spacing(
            p,
            {
                key: _style(element, key, None)
                for key in (
                    "line_spacing",
                    "line_spacing_pt",
                    "space_before_pt",
                    "space_after_pt",
                )
            },
        )
        align = str(
            _style(element, "align", _style(element, "horizontal_anchor", "left"))
        ).casefold()
        p.alignment = {
            "center": PP_ALIGN.CENTER,
            "middle": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
            "justify": PP_ALIGN.JUSTIFY,
        }.get(align, PP_ALIGN.LEFT)
        runs = _get(element, "runs", None) or _get(element, "text_runs", None)
        if runs:
            for spec in runs:
                for index, part in enumerate(
                    str(_get(spec, "text", "") or "").split("\n")
                ):
                    if index:
                        p.add_line_break()
                    run = p.add_run()
                    run.text = part
                    self._format_run(run, spec, element)
        else:
            for index, part in enumerate(str(text).split("\n")):
                if index:
                    p.add_line_break()
                run = p.add_run()
                run.text = part
                self._format_run(run, element, element)

    def _format_run(self, run, spec, element):
        family = _get(
            spec,
            "font_family",
            _style(element, "font_family", _style(element, "font", "Arial")),
        )
        size = float(
            _get(
                spec,
                "font_size_pt",
                _get(
                    spec,
                    "size",
                    _style(element, "font_size_pt", _style(element, "font_size", 12)),
                ),
            )
            or 12
        )
        weight = _get(
            spec,
            "weight",
            _get(
                spec,
                "font_weight",
                _style(element, "weight", _style(element, "font_weight", "regular")),
            ),
        )
        color = _get(
            spec,
            "color",
            _style(element, "color", _style(element, "font_color", "000000")),
        )
        set_run_typeface(run, str(family))
        run.font.size = __import__("pptx").util.Pt(size)
        run.font.bold = str(weight).casefold() in ("bold", "700", "800", "900") or bool(
            _get(spec, "bold", False)
        )
        run.font.italic = bool(_get(spec, "italic", _style(element, "italic", False)))
        run.font.color.rgb = _rgb(color)
        spacing = _get(
            spec,
            "char_spacing_pt",
            _style(element, "char_spacing_pt", None),
        )
        if spacing not in (None, 0, 0.0):
            rpr = run._r.get_or_add_rPr()
            rpr.set("spc", str(int(round(float(spacing) * 100))))

    @staticmethod
    def _set_alpha(xpr, opacity):
        # python-pptx has no public transparency API; inject alpha where possible.
        if not 0 <= opacity < 1:
            return
        try:
            ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
            # ``fill._xPr`` is the spPr parent.  ``find`` is not recursive, so
            # the srgbClr nested under solidFill was never reached and
            # fill_opacity compiled as an opaque color.
            solid = xpr.find(".//%ssrgbClr" % ns)
            if solid is None:
                return
            for child in list(solid):
                if child.tag == ns + "alpha":
                    solid.remove(child)
            solid.append(
                parse_xml(
                    '<a:alpha %s val="%d"/>' % (nsdecls("a"), int(round(opacity * 100000)))
                )
            )
        except Exception:
            pass

    def _gradient(self, shape, stops, angle):
        try:
            stops = normalize_gradient_stops(stops)
            sppr = shape._sp.spPr
            insert_at = None
            for index, child in enumerate(list(sppr)):
                if child.tag.rsplit("}", 1)[-1] in ("solidFill", "noFill", "gradFill"):
                    if insert_at is None:
                        insert_at = index
                    sppr.remove(child)
            def _stop(position, color):
                alpha = _split_alpha(color)[1]
                inner = "" if alpha is None or alpha >= 1 else '<a:alpha val="%d"/>' % int(round(max(0.0, alpha) * 100000))
                return '<a:gs pos="%d"><a:srgbClr val="%s">%s</a:srgbClr></a:gs>' % (
                    int(round(position * 100000)), _hex(color), inner)

            gs = "".join(_stop(position, color) for position, color in stops)
            grad = parse_xml(
                '<a:gradFill %s rotWithShape="1"><a:gsLst>%s</a:gsLst><a:lin ang="%d" scaled="0"/></a:gradFill>'
                % (nsdecls("a"), gs, int(round(float(angle) % 360 * 60000)))
            )
            if insert_at is None:
                # DrawingML order: xfrm, geometry, fill, ln, effects.
                names = [c.tag.rsplit("}", 1)[-1] for c in sppr]
                insert_at = next(
                    (i for i, n in enumerate(names) if n in ("ln", "effectLst", "effectDag", "scene3d", "sp3d", "extLst")),
                    len(names),
                )
            sppr.insert(insert_at, grad)
        except Exception as exc:
            self._warnings.append("gradient unavailable: %r" % (exc,))

    def _effects(self, shape, element, canvas):
        """Apply editable DrawingML shadow/glow effects from normalized style."""

        shadow = _style(element, "shadow", None)
        glow = _style(element, "glow", None)
        if not shadow and not glow:
            return
        try:
            sppr = shape._sp.spPr
            for child in list(sppr):
                if child.tag.rsplit("}", 1)[-1] == "effectLst":
                    sppr.remove(child)
            effects = []
            if isinstance(shadow, Mapping):
                color = _hex(shadow.get("color", "000000"))
                opacity = max(0.0, min(1.0, float(shadow.get("opacity", 0.25))))
                blur = max(
                    0, int(px_to_emu(float(shadow.get("blur_px", 6)), "x", canvas))
                )
                distance = max(
                    0, int(px_to_emu(float(shadow.get("distance_px", 3)), "x", canvas))
                )
                direction = int(float(shadow.get("angle", 45)) * 60000)
                effects.append(
                    '<a:outerShdw blurRad="%d" dist="%d" dir="%d" algn="ctr" rotWithShape="0">'
                    '<a:srgbClr val="%s"><a:alpha val="%d"/></a:srgbClr></a:outerShdw>'
                    % (blur, distance, direction, color, int(opacity * 100000))
                )
            if isinstance(glow, Mapping):
                color = _hex(glow.get("color", "FFFFFF"))
                opacity = max(0.0, min(1.0, float(glow.get("opacity", 0.4))))
                radius = max(
                    0, int(px_to_emu(float(glow.get("radius_px", 4)), "x", canvas))
                )
                effects.append(
                    '<a:glow rad="%d"><a:srgbClr val="%s"><a:alpha val="%d"/>'
                    "</a:srgbClr></a:glow>" % (radius, color, int(opacity * 100000))
                )
            if effects:
                sppr.append(
                    parse_xml(
                        "<a:effectLst %s>%s</a:effectLst>"
                        % (nsdecls("a"), "".join(effects))
                    )
                )
        except Exception as exc:
            self._warnings.append("effects unavailable: %s" % exc)

    @staticmethod
    def _apply_shape_adjust(shape, element):
        """Write preset-geometry adjustments.  Values are 0–1 fractions."""

        raw = _style(element, "shape_adjust", None)
        if raw is None:
            return
        values = raw if isinstance(raw, (list, tuple)) else (raw,)
        try:
            for index, value in enumerate(values):
                shape.adjustments[index] = float(value)
        except (AttributeError, IndexError, TypeError, ValueError):
            return

    @staticmethod
    def _set_corner_adjustment(shape, radius_px, width_emu, height_emu, canvas):
        try:
            # adjustment 0..0.5 is a proportion of the smaller dimension.
            value = min(
                0.5,
                max(
                    0.0,
                    px_to_emu(radius_px, "x", canvas)
                    / max(1.0, min(width_emu, height_emu)),
                ),
            )
            # AdjustmentCollection returns a float and exposes assignment via
            # __setitem__; it is not an Adjustment object with ``.value``.
            shape.adjustments[0] = value
        except Exception:
            pass


def _is_picture_shape(shape: Any) -> bool:
    """Recognize python-pptx's ``PICTURE (13)`` enum reliably."""

    value = getattr(shape, "shape_type", None)
    if value is None:
        return False
    name = getattr(value, "name", None)
    return str(name or value).upper().split(" ", 1)[0] == "PICTURE"


__all__ = ["BuildResult", "NativePrimitives", "normalize_gradient_stops", "strip_theme_style"]


def normalize_gradient_stops(stops):
    """Accept the stop shapes hosts actually write and return [(pos 0..1, color)].

    Supported: ``[pos, color]``, ``[color, pos]``, ``{"position"|"pos"|"offset":
    p, "color": c}`` and bare colours (spread evenly).  Positions may be 0..1 or
    0..100 (percent).
    """

    items = list(stops or [])
    if len(items) < 2:
        raise ValueError("gradient needs at least two stops")
    parsed = []
    for index, stop in enumerate(items):
        position = color = None
        if isinstance(stop, dict):
            color = stop.get("color", stop.get("colour", stop.get("hex")))
            for key in ("position", "pos", "offset", "at", "stop"):
                if key in stop:
                    position = stop[key]
                    break
            alpha = stop.get("opacity", stop.get("alpha"))
            if alpha is not None and isinstance(color, str):
                base, own = _split_alpha(color)
                a = float(alpha) * (1.0 if own is None else own)
                color = "#%s%02X" % (_hex(base), int(round(max(0.0, min(1.0, a)) * 255)))
        elif isinstance(stop, str):
            color = stop
        else:
            stop = list(stop)
            first, second = stop[:2]
            if isinstance(first, str) and not isinstance(second, str):
                color, position = first, second
            else:
                position, color = first, second
        if not isinstance(color, str):
            raise ValueError("gradient stop %d has no colour" % index)
        if isinstance(position, str):
            position = float(position.strip().rstrip("%")) / (100.0 if position.strip().endswith("%") else 1.0)
        parsed.append([None if position is None else float(position), color])
    if all(p is None for p, _ in parsed):
        last = len(parsed) - 1
        parsed = [[i / float(last), c] for i, (_, c) in enumerate(parsed)]
    if any(p is None for p, _ in parsed):
        raise ValueError("gradient stops mix positioned and unpositioned entries")
    if max(p for p, _ in parsed) > 1.0:
        parsed = [[p / 100.0, c] for p, c in parsed]
    return [(min(1.0, max(0.0, p)), c) for p, c in parsed]
