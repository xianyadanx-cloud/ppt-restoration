"""Content inventory and deterministic PPT DOM reconciliation.

The inventory is generated from the *audited SceneSpec*, not from OCR.  It is
therefore a preservation check for ``model output -> PPT DOM``.  It must never
be presented as proof that the model recognized the source image correctly.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Union

from ppt_restore.contracts.schema_v2 import SceneNode, SceneSpecV2, scene_sha256

PathLike = Union[str, Path]


def _text(value: Any) -> str:
    return str(value if value is not None else "").strip()


def _run_text(runs: Any) -> str:
    if not isinstance(runs, (list, tuple)):
        return ""
    return "".join(
        _text(run.get("text", "")) for run in runs if isinstance(run, Mapping)
    ).strip()


def _cell_text(cell: Any) -> str:
    if isinstance(cell, Mapping):
        return _text(cell.get("text", cell.get("value", "")))
    return _text(cell)


def _visible_step(step: Any) -> Any:
    if isinstance(step, Mapping):
        for key in ("text", "label", "name", "title", "value"):
            if key in step:
                return step[key]
        return dict(step)
    return step


def node_content(node: SceneNode) -> Dict[str, Any]:
    """Return the content-bearing fields that should survive compilation."""

    payload = dict(node.payload or {})
    kind = str(node.kind)
    if kind == "text":
        value = payload.get("text")
        if value is None:
            value = _run_text(payload.get("runs"))
        return {"text": str(value or "")}
    if kind == "badge":
        return {"text": str(payload.get("text", payload.get("label", "")) or "")}
    if kind == "kpi_card":
        return {
            key: payload[key]
            for key in ("label", "value", "delta")
            if payload.get(key) is not None
        }
    if kind == "progress_bar":
        # The numerical fraction is geometry, while ``label`` is the visible
        # content.  Do not require a hidden numeric payload in the PPT DOM.
        return {"label": payload["label"]} if payload.get("label") is not None else {}
    if kind == "table":
        cells = payload.get("cells", [])
        if isinstance(cells, (list, tuple)):
            return {
                "table_cells": [
                    [_cell_text(cell) for cell in row]
                    if isinstance(row, (list, tuple))
                    else [_cell_text(row)]
                    for row in cells
                ]
            }
        return {"table_cells": []}
    if kind == "chart":
        content: Dict[str, Any] = {}
        if "categories" in payload:
            content["categories"] = list(payload.get("categories") or [])
        if "series" in payload:
            content["series"] = [
                {
                    "name": item.get("name", "Series"),
                    "values": list(item.get("values", []) or ()),
                }
                if isinstance(item, Mapping)
                else item
                for item in (payload.get("series") or ())
            ]
        return content
    if kind == "process":
        steps = payload.get("steps")
        return (
            {"steps": [_visible_step(step) for step in steps]}
            if isinstance(steps, (list, tuple))
            else {}
        )
    # Shape-like payloads can carry an explicit label/text.  Pure containers,
    # lines, and decoration geometry intentionally contribute no content.
    for key in ("text", "label", "value"):
        if payload.get(key) is not None:
            return {key: payload[key]}
    return {}


def _bbox_norm(node: SceneNode, scene: SceneSpecV2) -> List[float]:
    width = float(scene.canvas.get("width_px", 0) or 0)
    height = float(scene.canvas.get("height_px", 0) or 0)
    x, y, w, h = node.bbox_px
    return [
        round(x * 1000.0 / width, 6),
        round(y * 1000.0 / height, 6),
        round(w * 1000.0 / width, 6),
        round(h * 1000.0 / height, 6),
    ]


def _stable_scene_sha256(scene: SceneSpecV2) -> str:
    normalized = SceneSpecV2.from_dict(json.loads(scene.to_json()))
    return scene_sha256(normalized)


def _audit_items(audit: Any) -> Mapping[str, Mapping[str, Any]]:
    if audit is None:
        return {}
    raw = (
        audit.get("items", ())
        if isinstance(audit, Mapping)
        else getattr(audit, "items", ())
    )
    result = {}
    for item in raw or ():
        if isinstance(item, Mapping) and item.get("node_id"):
            result[str(item["node_id"])] = dict(item)
    return result


def scene_to_content_inventory(
    scene: SceneSpecV2,
    audit: Any = None,
    *,
    audit_sha256: Optional[str] = None,
) -> Dict[str, Any]:
    """Create a stable content inventory from a SceneSpec and optional audit."""

    audit_items = _audit_items(audit)
    audit_status = "NOT_PROVIDED_COMPATIBILITY"
    if audit is not None:
        audit_status = str(
            audit.get("status")
            if isinstance(audit, Mapping)
            else getattr(audit, "status", "PROPOSED")
        )
    items = []
    for node in scene.nodes:
        item = audit_items.get(node.id, {})
        content = node_content(node)
        items.append(
            {
                "node_id": node.id,
                "role": str(node.role),
                "kind": str(node.kind),
                "content": content,
                "bbox_norm": _bbox_norm(node, scene),
                "confidence": float(item.get("confidence", node.confidence)),
                "evidence_refs": list(item.get("evidence_refs", node.evidence_refs)),
                "critical": bool(
                    node.role not in ("decoration", "container")
                    and node.kind != "group"
                ),
                "audit_item_status": str(item.get("status", "NOT_PROVIDED")),
            }
        )
    inventory: Dict[str, Any] = {
        "schema_version": "1.0",
        "scene_schema_version": str(scene.schema_version),
        "scene_sha256": _stable_scene_sha256(scene),
        "canonical_sha256": str(scene.canonical_sha256),
        "evidence_sha256": str(scene.evidence_sha256),
        "audit_sha256": str(audit_sha256) if audit_sha256 else None,
        "audit_status": audit_status,
        "items": items,
        "claim": "proves model-submitted SceneSpec content was preserved in the PPT DOM",
        "not_source_recognition_proof": True,
    }
    return inventory


def write_content_inventory(path: PathLike, inventory: Mapping[str, Any]) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        json.dumps(dict(inventory), ensure_ascii=False, indent=2, sort_keys=True)
        + "\n",
        encoding="utf-8",
    )
    return target


def read_content_inventory(path: PathLike) -> Dict[str, Any]:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(data, Mapping):
        raise ValueError("content inventory must be a JSON object")
    for key in (
        "schema_version",
        "scene_sha256",
        "canonical_sha256",
        "evidence_sha256",
        "items",
    ):
        if key not in data:
            raise ValueError("content inventory missing %s" % key)
    if str(data["schema_version"]) != "1.0":
        raise ValueError("unsupported content inventory schema version")
    if not isinstance(data["items"], list):
        raise ValueError("content inventory items must be a list")
    return dict(data)


def _walk_shapes(shapes: Any) -> Iterable[Any]:
    for shape in shapes or ():
        yield shape
        nested = getattr(shape, "shapes", None)
        if nested is not None:
            yield from _walk_shapes(nested)


def _dom_content(pptx_path: PathLike) -> Dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - dependency is core
        raise RuntimeError("python-pptx is required for DOM reconciliation") from exc
    presentation = Presentation(str(pptx_path))
    shape_texts: List[str] = []
    table_cells: List[str] = []
    chart_texts: List[str] = []
    for slide in presentation.slides:
        for shape in _walk_shapes(slide.shapes):
            try:
                if bool(getattr(shape, "has_table", False)):
                    for row in shape.table.rows:
                        for cell in row.cells:
                            value = str(cell.text or "").strip()
                            if value:
                                table_cells.append(value)
                    continue
                if bool(getattr(shape, "has_text_frame", False)):
                    value = str(shape.text or "").strip()
                    if value:
                        shape_texts.append(value)
                if bool(getattr(shape, "has_chart", False)):
                    chart = shape.chart
                    if bool(getattr(chart, "has_title", False)):
                        value = str(chart.chart_title.text_frame.text or "").strip()
                        if value:
                            chart_texts.append(value)
            except Exception:
                # A malformed optional chart/shape should not hide the text
                # already extracted from the rest of the DOM.
                continue
    return {
        "shape_texts": shape_texts,
        "table_cells": table_cells,
        "chart_texts": chart_texts,
        "all_text": shape_texts + table_cells + chart_texts,
    }


def _flatten_content(value: Any) -> Iterable[str]:
    if isinstance(value, Mapping):
        for child in value.values():
            yield from _flatten_content(child)
    elif isinstance(value, (list, tuple)):
        for child in value:
            yield from _flatten_content(child)
    elif value is not None:
        text = str(value).strip()
        if text:
            yield text


def _norm_text(value: Any) -> str:
    return " ".join(str(value or "").split())


def _is_present(expected: str, dom_values: Sequence[str]) -> bool:
    target = _norm_text(expected)
    if not target:
        return True
    normalized = [_norm_text(value) for value in dom_values if _norm_text(value)]
    if target in normalized:
        return True
    # Runs and table cells may be combined by python-pptx; use a conservative
    # containment fallback for longer strings but avoid making "1" match "10".
    if len(target) >= 3 and target in "\n".join(normalized):
        return True
    return False


def reconcile_inventory_with_pptx(
    inventory: Union[PathLike, Mapping[str, Any]],
    pptx_path: PathLike,
    *,
    report_path: Optional[PathLike] = None,
) -> Dict[str, Any]:
    """Compare every inventory content value with text/table content in PPTX."""

    inventory_path = Path(inventory) if isinstance(inventory, (str, Path)) else None
    data = (
        read_content_inventory(inventory)
        if inventory_path is not None
        else dict(inventory)
    )  # type: ignore[arg-type]
    dom = _dom_content(pptx_path)
    missing_critical = []
    missing = []
    for item in data.get("items", ()):
        if not isinstance(item, Mapping):
            continue
        expected = list(_flatten_content(item.get("content", {})))
        absent = [
            value for value in expected if not _is_present(value, dom["all_text"])
        ]
        if not absent:
            continue
        entry = {
            "node_id": str(item.get("node_id", "")),
            "missing_content": absent,
            "critical": bool(item.get("critical", False)),
        }
        missing.append(entry)
        if entry["critical"]:
            missing_critical.append(entry)
    valid = not missing_critical
    report: Dict[str, Any] = {
        "schema_version": "1.0",
        "status": "PASS" if valid else "NEEDS_REVIEW",
        "valid": valid,
        "inventory_path": str(inventory_path.resolve()) if inventory_path else None,
        "pptx_path": str(Path(pptx_path).resolve()),
        "missing": missing,
        "missing_critical": missing_critical,
        "dom": dom,
        "claim": "deterministic model-output-to-PPT-DOM preservation check; not proof of source-image recognition",
    }
    if inventory_path is not None:
        report["inventory_sha256"] = hashlib.sha256(
            inventory_path.read_bytes()
        ).hexdigest()
    pptx = Path(pptx_path)
    if pptx.is_file():
        report["pptx_sha256"] = hashlib.sha256(pptx.read_bytes()).hexdigest()
    if report_path is not None:
        target = Path(report_path)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(
            json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
    return report


__all__ = [
    "node_content",
    "scene_to_content_inventory",
    "write_content_inventory",
    "read_content_inventory",
    "reconcile_inventory_with_pptx",
]
